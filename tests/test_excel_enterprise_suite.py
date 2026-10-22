# -*- coding: utf-8 -*-
"""
test_excel_enterprise_suite.py — Automated Test Suite for Enterprise Excel Data Suite.

Verifies:
- ERR_XLSX_009: 15-Digit Precision & Accounting Negative Guard (CCCD, card numbers, Vietnamese dates, (100k) -> -100k, '-' -> 0, localized decimals).
- ERR_XLSX_010: XML Illegal Character Sanitization & 32k Limit.
- ERR_XLSX_011: _xlfn Modern Formula Normalization (Dynamic Array & Modern formulas).
- ERR_XLSX_012: Gridlines, AutoFilter & Unlocked Row Height.
- ERR_XLSX_013: Selective Formula & Structure Protection.
- ERR_XLSX_014: RAM-Constrained Big Data Streaming Mode (<50MB RAM).
- Subsystem 1: Automated Conditional Formatting (Status Highlights, Data Bars, Color Scales).
- Subsystem 2: Data Validation & Dropdowns (List, Length, Date).
- Subsystem 3: Native In-Sheet Charts (Bar/Line/Pie with ERR_DIAG_003 safe anchoring).
- Subsystem 6: Executive Index Dashboard & 2-Way Hyperlink Navigation.
"""

from __future__ import annotations

import datetime
import os
import sys
import tempfile
import unittest

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from core.xlsx import (
    UniversalExcelEngine,
    safe_set_cell,
    sanitize_xml_string,
    parse_vietnamese_date,
    parse_accounting_and_localized_number,
    normalize_modern_formula,
    stream_read_xlsx,
    stream_write_xlsx,
    estimate_xlsx_metrics,
)


class TestExcelEnterpriseSuite(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dpath = self.temp_dir.name

    def tearDown(self):
        self.temp_dir.cleanup()

    # -----------------------------------------------------------------------
    # Edge Case 1: 15-Digit Precision & Accounting Negative Guard (ERR_XLSX_009)
    # -----------------------------------------------------------------------

    def test_15_digit_precision_and_cccd_guard(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None

        # CCCD 12 digits (with leading zero)
        safe_set_cell(ws, "A1", value="001200123456")
        self.assertEqual(ws["A1"].number_format, "@")
        self.assertEqual(ws["A1"].value, "001200123456")
        self.assertIsInstance(ws["A1"].value, str)

        # Credit card 16 digits
        safe_set_cell(ws, "A2", value="9704123456789012")
        self.assertEqual(ws["A2"].number_format, "@")
        self.assertEqual(ws["A2"].value, "9704123456789012")

        # Bank account 19 digits
        safe_set_cell(ws, "A3", value="1234567890123456789")
        self.assertEqual(ws["A3"].number_format, "@")
        self.assertEqual(ws["A3"].value, "1234567890123456789")

    def test_vietnamese_date_parsing(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None

        # DD/MM/YYYY
        safe_set_cell(ws, "B1", value="26/05/2026")
        self.assertEqual(ws["B1"].value, datetime.date(2026, 5, 26))
        self.assertEqual(ws["B1"].number_format, "DD/MM/YYYY")

        # 'ngày 26 tháng 05 năm 2026'
        safe_set_cell(ws, "B2", value="ngày 26 tháng 05 năm 2026")
        self.assertEqual(ws["B2"].value, datetime.date(2026, 5, 26))
        self.assertEqual(ws["B2"].number_format, "DD/MM/YYYY")

        # Fallback for invalid date (e.g. 31/02/2026) -> remains string
        safe_set_cell(ws, "B3", value="31/02/2026")
        self.assertEqual(ws["B3"].value, "31/02/2026")

    def test_accounting_negatives_and_dash(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None

        # Dash '-' -> 0 with accounting format
        safe_set_cell(ws, "C1", value="-")
        self.assertEqual(ws["C1"].value, 0)
        self.assertIn("#,##0", ws["C1"].number_format)

        # Accounting negative '(100.000)' -> -100000
        safe_set_cell(ws, "C2", value="(100.000)")
        self.assertEqual(ws["C2"].value, -100000)
        self.assertIn("#,##0", ws["C2"].number_format)

        # Accounting negative with decimal '(20.000,50)' -> -20000.50
        safe_set_cell(ws, "C3", value="(20.000,50)")
        self.assertEqual(ws["C3"].value, -20000.50)
        self.assertIn("#,##0.00", ws["C3"].number_format)

    def test_vietnamese_locale_decimals(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None

        # '20.000,50' -> 20000.50
        safe_set_cell(ws, "D1", value="20.000,50")
        self.assertEqual(ws["D1"].value, 20000.50)

        # '1.250.000' -> 1250000
        safe_set_cell(ws, "D2", value="1.250.000")
        self.assertEqual(ws["D2"].value, 1250000)

    # -----------------------------------------------------------------------
    # Edge Case 2: XML Sanitization & 32k Limit (ERR_XLSX_010)
    # -----------------------------------------------------------------------

    def test_xml_sanitization_and_32k_limit(self):
        # Strip ASCII illegal characters \x00, \x08, \x0b, \x1f
        corrupt_str = "Tài liệu\x00 báo cáo\x08 tài chính\x0b quý 2\x1f."
        clean = sanitize_xml_string(corrupt_str)
        self.assertEqual(clean, "Tài liệu báo cáo tài chính quý 2.")

        # Truncate string > 32,000 characters
        huge_str = "A" * 35000
        truncated = sanitize_xml_string(huge_str)
        self.assertTrue(len(truncated) <= 32100)
        self.assertTrue(truncated.endswith("... [Xem tiếp tài liệu gốc]"))

    # -----------------------------------------------------------------------
    # Edge Case 3: Modern Formula _xlfn. Compatibility (ERR_XLSX_011)
    # -----------------------------------------------------------------------

    def test_modern_formula_normalization(self):
        # XLOOKUP
        f1 = normalize_modern_formula("=XLOOKUP(A1, B:B, C:C)")
        self.assertEqual(f1, "=_xlfn.XLOOKUP(A1, B:B, C:C)")

        # CONCAT & TEXTJOIN
        f2 = normalize_modern_formula('=CONCAT(A1, " - ", TEXTJOIN(", ", TRUE, B1:B10))')
        self.assertEqual(f2, '=_xlfn.CONCAT(A1, " - ", _xlfn.TEXTJOIN(", ", TRUE, B1:B10))')

        # Prevent double prefixing
        f3 = normalize_modern_formula("=_xlfn.XLOOKUP(A1, B:B, C:C)")
        self.assertEqual(f3, "=_xlfn.XLOOKUP(A1, B:B, C:C)")

        # In-flight formula via safe_set_cell
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        safe_set_cell(ws, "E1", formula="=UNIQUE(FILTER(A1:A10, B1:B10>0))")
        self.assertEqual(ws["E1"].value, "=_xlfn.UNIQUE(_xlfn.FILTER(A1:A10, B1:B10>0))")

    # -----------------------------------------------------------------------
    # Edge Case 4 & Subsystems: Gridlines, AutoFilter, Row Height (ERR_XLSX_012)
    # -----------------------------------------------------------------------

    def test_gridlines_and_autofilter(self):
        wb = openpyxl.Workbook()
        engine = UniversalExcelEngine(wb)
        ws = engine.get_sheet("Sheet")

        # Headers
        ws["A1"] = "ID"
        ws["B1"] = "Name"
        ws["C1"] = "Score"
        ws["A2"] = "001"
        ws["B2"] = "Alpha"
        ws["C2"] = 100

        engine.enable_gridlines("Sheet")
        engine.set_autofilter("Sheet", header_row=1)

        self.assertTrue(ws.views.sheetView[0].showGridLines)
        self.assertEqual(ws.auto_filter.ref, "A1:C2")

    # -----------------------------------------------------------------------
    # Subsystem 1: Automated Conditional Formatting
    # -----------------------------------------------------------------------

    def test_conditional_formatting_and_data_bars(self):
        wb = openpyxl.Workbook()
        engine = UniversalExcelEngine(wb)
        ws = engine.get_sheet("CF_Test")

        # Populate sample status and metric
        ws["A1"] = "Status"
        ws["A2"] = "Chưa duyệt"
        ws["A3"] = "Đã duyệt"
        ws["B1"] = "Progress"
        ws["B2"] = 25
        ws["B3"] = 80

        engine.apply_status_highlights("CF_Test", "A2:A10")
        engine.apply_data_bars("CF_Test", "B2:B10", color="3B82F6")
        engine.apply_color_scale("CF_Test", "B2:B10", scale_type="green_yellow_red")

        cf_list = list(ws.conditional_formatting)
        self.assertGreaterEqual(len(cf_list), 2)

    # -----------------------------------------------------------------------
    # Subsystem 2: Data Validation & Dropdowns
    # -----------------------------------------------------------------------

    def test_data_validation_subsystem(self):
        wb = openpyxl.Workbook()
        engine = UniversalExcelEngine(wb)
        ws = engine.get_sheet("DV_Test")

        # 1. Dropdown
        engine.add_dropdown_validation("DV_Test", "A2:A50", ["Hà Nội", "Đà Nẵng", "TP.HCM"])

        # 2. Length (CCCD 12 digits)
        engine.add_length_validation("DV_Test", "B2:B50", exact_len=12)

        # 3. Date 2026
        engine.add_date_validation("DV_Test", "C2:C50", start_date="2026-01-01", end_date="2026-12-31")

        dv_rules = ws.data_validations.dataValidation
        self.assertEqual(len(dv_rules), 3)

    # -----------------------------------------------------------------------
    # Subsystem 3: Native In-Sheet Charts (ERR_DIAG_003)
    # -----------------------------------------------------------------------

    def test_native_in_sheet_chart(self):
        wb = openpyxl.Workbook()
        engine = UniversalExcelEngine(wb)
        ws = engine.get_sheet("Chart_Test")

        # Build sample table
        ws["A1"] = "Kênh"
        ws["B1"] = "Doanh Số"
        channels = [("Quầy", 500), ("Online", 1200), ("Đối tác", 350)]
        for r_i, (ch, val) in enumerate(channels, start=2):
            ws.cell(row=r_i, column=1, value=ch)
            ws.cell(row=r_i, column=2, value=val)

        chart = engine.add_native_chart(
            "Chart_Test",
            chart_type="bar",
            data_range=(2, 1, 2, 4),
            category_range=(1, 2, 1, 4),
            title="Biểu Đồ Doanh Số Phân Kênh",
            target_cell="auto",
        )
        self.assertIsNotNone(chart)
        self.assertEqual(len(ws._charts), 1)

    # -----------------------------------------------------------------------
    # Subsystem 4: Selective Protection & Formula Locking (ERR_XLSX_013)
    # -----------------------------------------------------------------------

    def test_selective_protection(self):
        wb = openpyxl.Workbook()
        engine = UniversalExcelEngine(wb)
        ws = engine.get_sheet("Protected_Test")

        # Header
        ws["A1"] = "Số Tiền"
        ws["B1"] = "Thuế (10%)"
        # Data
        ws["A2"] = 1000
        ws["B2"] = "=A2*0.1"

        # Allow user to edit A2, lock formula in B2 and header in A1:B1
        engine.protect_sheet_selective("Protected_Test", input_ranges=["A2"], password="admin")

        self.assertTrue(ws.protection.sheet)
        self.assertFalse(ws["A2"].protection.locked)
        self.assertTrue(ws["B2"].protection.locked)
        self.assertTrue(ws["A1"].protection.locked)

    # -----------------------------------------------------------------------
    # Subsystem 5: Big Data Streaming Mode (ERR_XLSX_014)
    # -----------------------------------------------------------------------

    def test_big_data_streaming_mode(self):
        stream_path = os.path.join(self.dpath, "stream_test.xlsx")

        def row_generator():
            for i in range(1, 1001):
                yield [f"ID_{i:04d}", f"User {i}", i * 10, "2026-05-26"]

        # 1. Stream Write
        dest = stream_write_xlsx(stream_path, row_generator(), sheet_name="Transactions", headers=["ID", "Name", "Score", "Date"])
        self.assertTrue(os.path.isfile(dest))

        # 2. Metric Estimation
        metrics = estimate_xlsx_metrics(stream_path)
        self.assertGreater(metrics["file_size_bytes"], 0)
        self.assertIn("Transactions", metrics["sheet_names"])

        # 3. Stream Read
        read_rows = list(stream_read_xlsx(stream_path, sheet_name="Transactions", max_rows=10))
        self.assertEqual(len(read_rows), 10)
        self.assertEqual(read_rows[0], ["ID", "Name", "Score", "Date"])
        self.assertEqual(read_rows[1][0], "ID_0001")

    # -----------------------------------------------------------------------
    # Subsystem 6: Executive Index Dashboard & Navigation (ERR_XLSX_012)
    # -----------------------------------------------------------------------

    def test_index_dashboard_and_navigation(self):
        wb = openpyxl.Workbook()
        # Create 3 child sheets
        wb.create_sheet("BieuPhi")
        wb.create_sheet("DoanhThu")
        wb.create_sheet("BaoCaoQuy")
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        engine = UniversalExcelEngine(wb)
        idx_ws = engine.create_index_dashboard(
            title="Mục Lục & Bảng Biểu Tổng Hợp",
            descriptions={"BieuPhi": "Biểu phí tài khoản cá nhân", "DoanhThu": "Kê khai doanh thu dịch vụ"},
            back_links=True
        )

        self.assertEqual(wb.sheetnames[0], "Mục Lục")
        self.assertIn("=HYPERLINK", idx_ws["E4"].value)
        self.assertIn("BieuPhi", idx_ws["E4"].value)

        # Check back link on child sheet
        child_ws = wb["BieuPhi"]
        self.assertIn("=HYPERLINK", child_ws["A1"].value)
        self.assertIn("Mục Lục", child_ws["A1"].value)


if __name__ == "__main__":
    unittest.main()
