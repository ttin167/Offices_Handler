"""
Unit test suite for Semantic Table Extraction to Excel (.xlsx) - Universal 5-Way Matrix.
Tests core.converter.table_extractor:
- Filtering administrative layout tables (Decree 30 headers and signatures)
- OpenXML <w:gridSpan> (colspan) and <w:vMerge> (rowspan) handling without MergedCell error (ERR_XLSX_004)
- Identifier number format '@' preserving leading zeros (ERR_XLSX_007)
- Sheet title sanitization and deduplication
- Markdown GFM table extraction with heading sheet mapping
- Empty table guard returning user-friendly Vietnamese message
- Universal conversion routes in core.converter.engine
"""

from __future__ import annotations

import os
import tempfile
import unittest

import openpyxl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.converter.table_extractor import (
    _sanitize_sheet_title,
    _is_layout_header_or_signature_table,
    extract_docx_tables_to_xlsx,
    extract_md_tables_to_xlsx,
    extract_tables_to_xlsx,
)
from core.converter.engine import convert_universal


class TestUniversalTableExtractor(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dpath = self.temp_dir.name

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_sanitize_sheet_title(self):
        used_titles: set[str] = set()
        # Test illegal characters
        title1 = _sanitize_sheet_title("Bảng: 1 / 2 * [Test] ?", used_titles, 1)
        self.assertNotIn(":", title1)
        self.assertNotIn("/", title1)
        self.assertNotIn("?", title1)
        self.assertNotIn("*", title1)
        self.assertNotIn("[", title1)
        self.assertNotIn("]", title1)
        self.assertLessEqual(len(title1), 31)

        # Test deduplication
        used_titles.add(title1.lower())
        title2 = _sanitize_sheet_title("Bảng: 1 / 2 * [Test] ?", used_titles, 2)
        self.assertNotEqual(title1.lower(), title2.lower())
        self.assertTrue(title2.endswith("_2"))

    def test_administrative_layout_table_filtered(self):
        """Decree 30 National Header and Signature tables must be skipped."""
        doc = Document()

        # Table 1: Decree 30 National Header layout table
        tbl_hdr = doc.add_table(rows=2, cols=2)
        tbl_hdr.cell(0, 0).text = "BỘ TÀI CHÍNH\nCỤC QUẢN LÝ DỰ ÁN"
        tbl_hdr.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
        tbl_hdr.cell(1, 0).text = "Số: 123/QĐ-BTC"
        tbl_hdr.cell(1, 1).text = "Hà Nội, ngày 22 tháng 10 năm 2026"

        self.assertTrue(_is_layout_header_or_signature_table(tbl_hdr))

        # Table 2: Real business data table
        doc.add_paragraph("Bảng thống kê ngân sách", style="Heading 2")
        tbl_data = doc.add_table(rows=3, cols=3)
        headers = ["Mã Dự Án", "Tên Hạng Mục", "Ngân Sách (Tỷ)"]
        for i, h in enumerate(headers):
            tbl_data.cell(0, i).text = h
        tbl_data.cell(1, 0).text = "001"
        tbl_data.cell(1, 1).text = "Nâng cấp máy chủ"
        tbl_data.cell(1, 2).text = "15.5"
        tbl_data.cell(2, 0).text = "002"
        tbl_data.cell(2, 1).text = "Bảo trì mạng WAN"
        tbl_data.cell(2, 2).text = "8.2"

        self.assertFalse(_is_layout_header_or_signature_table(tbl_data))

        # Table 3: Signature block table
        tbl_sig = doc.add_table(rows=2, cols=2)
        tbl_sig.cell(0, 0).text = "Nơi nhận:\n- Như điều 3;\n- Lưu VT."
        tbl_sig.cell(0, 1).text = "GIÁM ĐỐC DỰ ÁN\n(Ký và ghi rõ họ tên)"

        self.assertTrue(_is_layout_header_or_signature_table(tbl_sig))

        docx_path = os.path.join(self.dpath, "admin_doc.docx")
        doc.save(docx_path)

        xlsx_path = os.path.join(self.dpath, "extracted.xlsx")
        ok, res_path, count = extract_docx_tables_to_xlsx(docx_path, xlsx_path)

        self.assertTrue(ok)
        self.assertEqual(count, 1)  # Only the data table should be extracted!

        wb = openpyxl.load_workbook(res_path)
        self.assertEqual(len(wb.sheetnames), 1)
        ws = wb.active
        # Check that headers and rows were extracted
        self.assertEqual(ws.cell(row=1, column=1).value, "Mã Dự Án")
        self.assertEqual(ws.cell(row=2, column=1).value, "001")
        self.assertEqual(ws.cell(row=2, column=1).number_format, "@")
        wb.close()

    def test_docx_merged_cells_gridspan_and_vmerge(self):
        """OpenXML gridSpan and vMerge should not crash openpyxl with MergedCell read-only error."""
        doc = Document()
        tbl = doc.add_table(rows=3, cols=3)

        # Merge row 0 cell 0 and cell 1 horizontally (gridSpan=2)
        tc_00 = tbl.cell(0, 0)._tc
        tcPr_00 = tc_00.get_or_add_tcPr()
        gridSpan_00 = OxmlElement("w:gridSpan")
        gridSpan_00.set(qn("w:val"), "2")
        tcPr_00.append(gridSpan_00)
        tbl.cell(0, 0).text = "Tiêu Đề Gộp Ngang"
        tbl.cell(0, 2).text = "Cột Đơn"

        # Row 1 and Row 2 Cell 0 merged vertically (vMerge: restart then continue)
        tc_10 = tbl.cell(1, 0)._tc
        tcPr_10 = tc_10.get_or_add_tcPr()
        vMerge_restart = OxmlElement("w:vMerge")
        vMerge_restart.set(qn("w:val"), "restart")
        tcPr_10.append(vMerge_restart)
        tbl.cell(1, 0).text = "Dữ Liệu Gộp Dọc"

        tc_20 = tbl.cell(2, 0)._tc
        tcPr_20 = tc_20.get_or_add_tcPr()
        vMerge_cont = OxmlElement("w:vMerge")
        vMerge_cont.set(qn("w:val"), "continue")
        tcPr_20.append(vMerge_cont)

        tbl.cell(1, 1).text = "D1"
        tbl.cell(1, 2).text = "D2"
        tbl.cell(2, 1).text = "D3"
        tbl.cell(2, 2).text = "D4"

        docx_path = os.path.join(self.dpath, "merged_table.docx")
        doc.save(docx_path)

        xlsx_path = os.path.join(self.dpath, "merged_table.xlsx")
        ok, res_path, count = extract_docx_tables_to_xlsx(docx_path, xlsx_path)

        self.assertTrue(ok)
        self.assertEqual(count, 1)

        wb = openpyxl.load_workbook(res_path)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Tiêu Đề Gộp Ngang")
        wb.close()

    def test_markdown_gfm_table_extraction(self):
        md_content = """# 1. Báo Cáo Doanh Thu 2026

Dưới đây là bảng số liệu tài chính quý 1:

| Quý | Doanh Thu (Tỷ) | Tăng Trưởng | Mã Dự Án |
| :--- | :--- | :--- | :--- |
| Q1 | 120.5 | 15.2% | 0092 |
| Q2 | 145.0 | 20.3% | 0093 |

# 2. Phân Hệ Nhân Sự

| Phòng Ban | Số Lượng Nhân Sự |
| :--- | :--- |
| Kỹ Thuật | 45 |
| Kinh Doanh | 30 |
"""
        md_path = os.path.join(self.dpath, "report.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)

        xlsx_path = os.path.join(self.dpath, "from_md.xlsx")
        ok, res_path, count = extract_md_tables_to_xlsx(md_path, xlsx_path)

        self.assertTrue(ok)
        self.assertEqual(count, 2)

        wb = openpyxl.load_workbook(res_path)
        self.assertEqual(len(wb.sheetnames), 2)
        # Check sheet 1
        ws1 = wb[wb.sheetnames[0]]
        self.assertEqual(ws1.cell(row=1, column=1).value, "Quý")
        self.assertEqual(ws1.cell(row=2, column=4).value, "0092")
        self.assertEqual(ws1.cell(row=2, column=4).number_format, "@")
        wb.close()

    def test_empty_table_guard(self):
        """Document with no tables should return clear error message."""
        doc = Document()
        doc.add_paragraph("Tài liệu chỉ có văn bản thông thường, không có bảng biểu nào.")
        docx_path = os.path.join(self.dpath, "no_tables.docx")
        doc.save(docx_path)

        xlsx_path = os.path.join(self.dpath, "no_tables.xlsx")
        ok, msg, count = extract_tables_to_xlsx(docx_path, xlsx_path)

        self.assertFalse(ok)
        self.assertIn("không chứa cấu trúc bảng biểu", msg)
        self.assertEqual(count, 0)

    def test_universal_converter_route_11_docx_to_xlsx(self):
        """Test convert_universal DOCX -> XLSX route."""
        doc = Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Khoản Mục"
        tbl.cell(0, 1).text = "Số Tiền"
        tbl.cell(1, 0).text = "Lương"
        tbl.cell(1, 1).text = "50000000"

        docx_path = os.path.join(self.dpath, "salary.docx")
        doc.save(docx_path)

        xlsx_path = os.path.join(self.dpath, "salary.xlsx")
        success, out_path, elapsed = convert_universal(docx_path, xlsx_path)

        self.assertTrue(success)
        self.assertTrue(os.path.exists(out_path))
        self.assertGreater(elapsed, 0.0)


if __name__ == "__main__":
    unittest.main()
