"""
test_xlsx_print_setup.py — Unit Tests for ERR_XLSX_008 Print-Ready Page Setup.

Verifies:
1. apply_worksheet_print_setup:
   - Portrait orientation when column count <= 7.
   - Landscape orientation when column count > 7.
   - fitToPage=True, fitToWidth=1, fitToHeight=0 (no horizontal page splitting).
   - Horizontal centering and margin standardization.
   - Auto-detection and setting of print_title_rows for header row repeating.
2. UniversalExcelEngine.apply_print_setup file-level in-place mutation.
3. Unified QA Diagnostic detection and non-destructive auto-repair of ERR_XLSX_008.
"""

import os
import sys
import unittest
import openpyxl
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from core.xlsx.writer import (
    apply_worksheet_print_setup,
    UniversalExcelEngine,
)
from tools.unified_qa_diagnostic import ExcelDiagnosticEngine


class TestXlsxPrintSetup(unittest.TestCase):
    def setUp(self):
        self.scratch_dir = os.path.join(BASE_DIR, "scratch", "test_xlsx_print")
        os.makedirs(self.scratch_dir, exist_ok=True)

    def test_apply_worksheet_print_setup_portrait(self):
        """Test sheet with <= 7 columns defaults to Portrait orientation and fit-to-width."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"

        # Populate 5 columns
        headers = ["ID", "Name", "Department", "Role", "Status"]
        ws.append(headers)
        for i in range(10):
            ws.append([i + 1, f"User {i}", "Engineering", "Developer", "Active"])

        apply_worksheet_print_setup(ws)

        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(ws.page_setup.fitToPage)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(ws.page_setup.fitToHeight, 0)
        self.assertEqual(ws.page_setup.orientation, ws.ORIENTATION_PORTRAIT)
        self.assertTrue(ws.print_options.horizontalCentered)
        self.assertEqual(ws.page_margins.left, 0.7)
        self.assertEqual(ws.page_margins.right, 0.7)

    def test_apply_worksheet_print_setup_landscape(self):
        """Test sheet with > 7 columns auto-switches to Landscape orientation."""
        wb = Workbook()
        ws = wb.active
        ws.title = "WideReport"

        # Populate 10 columns
        headers = [f"Col_{i}" for i in range(1, 11)]
        ws.append(headers)
        for r in range(5):
            ws.append([f"val_{r}_{c}" for c in range(1, 11)])

        apply_worksheet_print_setup(ws)

        self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
        self.assertTrue(ws.page_setup.fitToPage)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(ws.page_setup.orientation, ws.ORIENTATION_LANDSCAPE)

    def test_apply_worksheet_print_setup_header_repeating(self):
        """Test auto-discovery of header row for ws.print_title_rows."""
        wb = Workbook()
        ws = wb.active
        ws.title = "DataSheet"

        # Row 1 empty, Row 2 is table header
        ws.cell(row=2, column=1, value="Mã Nhân Viên")
        ws.cell(row=2, column=2, value="Họ và Tên")
        ws.cell(row=2, column=3, value="Email")

        for r in range(3, 20):
            ws.cell(row=r, column=1, value=f"NV{r:03d}")
            ws.cell(row=r, column=2, value=f"Nguyen Van {r}")
            ws.cell(row=r, column=3, value=f"user{r}@example.com")

        apply_worksheet_print_setup(ws)

        self.assertEqual(ws.print_title_rows, "$2:$2")

    def test_universal_excel_engine_apply_print_setup_file(self):
        """Test file-level apply_print_setup mutates all visible worksheets."""
        test_file = os.path.join(self.scratch_dir, "test_file_print_setup.xlsx")
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Tab1"
        ws1.append(["A", "B", "C"])
        ws1.append([1, 2, 3])

        ws2 = wb.create_sheet("Tab2")
        ws2.append([f"C{i}" for i in range(12)])
        ws2.append([i for i in range(12)])

        wb.save(test_file)

        # Mutate file using engine
        res = UniversalExcelEngine.apply_print_setup(test_file)
        self.assertTrue(res["success"])
        self.assertEqual(len(res["sheets_modified"]), 2)

        # Reopen and verify
        wb_checked = openpyxl.load_workbook(test_file)
        tab1 = wb_checked["Tab1"]
        tab2 = wb_checked["Tab2"]

        self.assertTrue(tab1.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(tab1.page_setup.orientation, "portrait")

        self.assertTrue(tab2.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(tab2.page_setup.orientation, "landscape")
        wb_checked.close()

    def test_unified_qa_diagnostic_err_xlsx_008_detect_and_repair(self):
        """Test Unified QA Diagnostic detects missing print setup and auto-repairs it."""
        unconfigured_file = os.path.join(self.scratch_dir, "unconfigured_report.xlsx")
        repaired_file = os.path.join(self.scratch_dir, "repaired_report.xlsx")

        wb = Workbook()
        ws = wb.active
        ws.title = "Financials"
        ws.append(["Mã", "Tên", "Doanh thu", "Chi phí", "Lợi nhuận"])
        for i in range(10):
            ws.append([f"P{i:02d}", f"Item {i}", 1000 + i * 50, 600, 400 + i * 50])
        wb.save(unconfigured_file)

        # 1. Diagnose: Expect ERR_XLSX_008
        issues, _ = ExcelDiagnosticEngine.diagnose(unconfigured_file)
        issue_codes = [issue.code for issue in issues]
        self.assertIn("ERR_XLSX_008", issue_codes)

        # 2. Repair
        ok, actions = ExcelDiagnosticEngine.repair(unconfigured_file, repaired_file)
        self.assertTrue(ok)
        self.assertTrue(any("ERR_XLSX_008" in act for act in actions))

        # 3. Diagnose repaired file: Expect ERR_XLSX_008 resolved
        repaired_issues, _ = ExcelDiagnosticEngine.diagnose(repaired_file)
        repaired_codes = [issue.code for issue in repaired_issues]
        self.assertNotIn("ERR_XLSX_008", repaired_codes)


if __name__ == "__main__":
    unittest.main()
