"""
tests.test_pdf_table_matrix_stitching — Unit & Integration Test Suite for PDF Table Matrix Processing.

Tests:
1. Ghost Column Pruning (_prune_ghost_columns)
2. 2-Tier Header Detection & Merged Ranges Calculation (_detect_table_header_structure)
3. 1-Tier Header Fallback (_detect_table_header_structure)
4. Multi-Page Continuous Table Stitching (_stitch_multipage_tables)
5. Multi-Page Distinct Tables Preservation (_stitch_multipage_tables)
6. End-to-End Extraction with Real Banking Fee Schedule PDF (106 rows, 9 cols, freeze_panes A3, 5 merged ranges)
7. PDF Toolkit Convenience Wrapper Integration (core.converter.pdf_toolkit)
"""

from __future__ import annotations

import os
import tempfile
import unittest

import openpyxl

from core.converter.table_extractor import (
    _prune_ghost_columns,
    _detect_table_header_structure,
    _normalize_row_for_comparison,
    _stitch_multipage_tables,
    _derive_pdf_sheet_title,
    extract_pdf_tables_to_xlsx,
)
from core.converter.pdf_toolkit import extract_pdf_tables_to_xlsx as toolkit_extract_pdf_tables


class TestPdfTableMatrixStitching(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dpath = self.temp_dir.name
        self.real_pdf_path = os.path.abspath(
            os.path.join(
                os.path.dirname(__file__),
                "..",
                ".web_uploads",
                "20260526_Bieu phi dich vu tai khoan_ban tieng Viet.pdf",
            )
        )

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_prune_ghost_columns(self):
        """Ghost columns with all empty or whitespace cells must be pruned."""
        matrix = [
            ["Col1", "", "Col2", None, "Col3"],
            ["Val1", "", "Val2", "   ", "Val3"],
            ["Val4", "", "Val5", "", "Val6"],
        ]
        pruned = _prune_ghost_columns(matrix)
        self.assertEqual(len(pruned), 3)
        self.assertEqual(len(pruned[0]), 3)
        self.assertEqual(pruned[0], ["Col1", "Col2", "Col3"])
        self.assertEqual(pruned[1], ["Val1", "Val2", "Val3"])
        self.assertEqual(pruned[2], ["Val4", "Val5", "Val6"])

    def test_detect_two_tier_header(self):
        """Accurately detects 2-tier header and outputs 5 merge ranges."""
        rows = [
            ["STT", "Mã phí", "Dịch vụ", "Kênh quầy", None, None, "Kênh NHĐT", None, None],
            [None, None, None, "Mức phí", "Tối thiểu", "Tối đa", "Mức phí", "Tối thiểu", "Tối đa"],
            ["I", "", "Mở tài khoản", "", "", "", "", "", ""],
            ["1", "A043", "Tài khoản thường", "Miễn phí", "", "", "Miễn phí", "", ""],
        ]
        depth, merges = _detect_table_header_structure(rows)
        self.assertEqual(depth, 2)
        # Expected: A1:A2, B1:B2, C1:C2 (vertical); D1:F1, G1:I1 (horizontal)
        expected_merges = {
            (1, 1, 2, 1),  # A1:A2
            (1, 2, 2, 2),  # B1:B2
            (1, 3, 2, 3),  # C1:C2
            (1, 4, 1, 6),  # D1:F1
            (1, 7, 1, 9),  # G1:I1
        }
        self.assertEqual(set(merges), expected_merges)

    def test_detect_single_tier_header(self):
        """Standard 1-tier header table returns depth 1 and no merge ranges."""
        rows = [
            ["STT", "Họ tên", "Số tiền"],
            ["1", "Nguyễn Văn A", "5000000"],
            ["2", "Trần Thị B", "8000000"],
        ]
        depth, merges = _detect_table_header_structure(rows)
        self.assertEqual(depth, 1)
        self.assertEqual(merges, [])

    def test_stitch_multipage_continuous_tables(self):
        """Tables across pages with matching headers are seamlessly stitched."""
        page_tables = [
            {
                "page": 0,
                "table_idx": 1,
                "rows": [
                    ["STT", "Dịch vụ", "Kênh quầy", None],
                    [None, None, "Mức", "Tối đa"],
                    ["1", "Gói 1", "Miễn phí", ""],
                ],
            },
            {
                "page": 1,
                "table_idx": 1,
                "rows": [
                    ["STT", "Dịch vụ", "Kênh quầy", None],
                    [None, None, "Mức", "Tối đa"],
                    ["2", "Gói 2", "100.000", "500.000"],
                    ["3", "Gói 3", "200.000", "1.000.000"],
                ],
            },
        ]
        stitched = _stitch_multipage_tables(page_tables)
        self.assertEqual(len(stitched), 1)
        group = stitched[0]
        # Total rows = 2 header + 1 (page 0 data) + 2 (page 1 data) = 5 rows
        self.assertEqual(len(group["rows"]), 5)
        self.assertEqual(group["rows"][0][0], "STT")
        self.assertEqual(group["rows"][2][0], "1")
        self.assertEqual(group["rows"][3][0], "2")
        self.assertEqual(group["rows"][4][0], "3")

    def test_stitch_independent_tables_preserved(self):
        """Different tables with differing column schemas must not be merged."""
        page_tables = [
            {
                "page": 0,
                "table_idx": 1,
                "rows": [
                    ["Col A", "Col B"],
                    ["1", "2"],
                ],
            },
            {
                "page": 1,
                "table_idx": 1,
                "rows": [
                    ["ID", "Name", "Score", "Rank"],
                    ["101", "Alice", "95", "1"],
                ],
            },
        ]
        stitched = _stitch_multipage_tables(page_tables)
        self.assertEqual(len(stitched), 2)
        self.assertEqual(len(stitched[0]["rows"]), 2)
        self.assertEqual(len(stitched[1]["rows"]), 2)

    def test_pdf_to_xlsx_end_to_end_real_document(self):
        """Tests extraction on the actual Vietcombank fee schedule PDF."""
        if not os.path.isfile(self.real_pdf_path):
            self.skipTest(f"Real sample PDF not found at {self.real_pdf_path}")

        out_xlsx = os.path.join(self.dpath, "bieu_phi_converted.xlsx")
        ok, res_path, count = extract_pdf_tables_to_xlsx(self.real_pdf_path, out_xlsx)

        self.assertTrue(ok)
        self.assertEqual(count, 1)
        self.assertTrue(os.path.isfile(out_xlsx))

        wb = openpyxl.load_workbook(out_xlsx)
        self.assertEqual(len(wb.sheetnames), 1)

        ws = wb.active
        self.assertEqual(ws.title, "Bieu_Phi_Dich_Vu")
        self.assertEqual(ws.max_row, 106)
        self.assertEqual(ws.max_column, 9)
        self.assertEqual(ws.freeze_panes, "A3")

        # Check 5 Merged cell ranges
        merge_strings = {str(r) for r in ws.merged_cells.ranges}
        expected_ranges = {"A1:A2", "B1:B2", "C1:C2", "D1:F1", "G1:I1"}
        self.assertEqual(merge_strings, expected_ranges)

        # Check Row 1 and Row 2 headers
        self.assertEqual(ws.cell(row=1, column=1).value, "STT")
        self.assertEqual(ws.cell(row=1, column=2).value, "Mã phí")
        self.assertEqual(ws.cell(row=1, column=3).value, "Dịch vụ")
        self.assertEqual(ws.cell(row=1, column=4).value, "Kênh quầy")
        self.assertEqual(ws.cell(row=1, column=7).value, "Kênh NHĐT")

        self.assertEqual(ws.cell(row=2, column=4).value, "Mức phí")
        self.assertEqual(ws.cell(row=2, column=5).value, "Tối thiểu")
        self.assertEqual(ws.cell(row=2, column=6).value, "Tối đa")
        self.assertEqual(ws.cell(row=2, column=7).value, "Mức phí")
        self.assertEqual(ws.cell(row=2, column=8).value, "Tối thiểu")
        self.assertEqual(ws.cell(row=2, column=9).value, "Tối đa")

        # Check Row 4 (A043)
        self.assertEqual(ws.cell(row=4, column=1).value, 1)
        self.assertEqual(ws.cell(row=4, column=2).value, "A043")
        self.assertEqual(ws.cell(row=4, column=3).value, "Tài khoản thường")
        self.assertEqual(ws.cell(row=4, column=4).value, "Miễn phí")
        self.assertEqual(ws.cell(row=4, column=7).value, "Miễn phí")

        # Check Row 106 (A041 - Last data row)
        self.assertEqual(ws.cell(row=106, column=1).value, 6)
        self.assertEqual(ws.cell(row=106, column=2).value, "A041")
        self.assertEqual(ws.cell(row=106, column=3).value, "Dịch vụ tài khoản khác")

        wb.close()

    def test_pdf_toolkit_wrapper(self):
        """Tests core.converter.pdf_toolkit.extract_pdf_tables_to_xlsx wrapper."""
        if not os.path.isfile(self.real_pdf_path):
            self.skipTest(f"Real sample PDF not found at {self.real_pdf_path}")

        out_xlsx = os.path.join(self.dpath, "toolkit_out.xlsx")
        res = toolkit_extract_pdf_tables(self.real_pdf_path, out_xlsx)

        self.assertTrue(res.get("success"))
        self.assertEqual(res.get("tables_extracted"), 1)
        self.assertTrue(os.path.isfile(res.get("output_path", "")))


if __name__ == "__main__":
    unittest.main()
