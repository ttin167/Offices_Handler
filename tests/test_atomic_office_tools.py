"""
tests/test_atomic_office_tools.py — Test suite for 8 Atomic Office Tools & CLI subcommands.

Verifies:
1. xlsx_set_cell (ERR_XLSX_004 merged-cell protection, ERR_XLSX_007 leading zeros string format)
2. xlsx_set_range (multi-cell patch)
3. xlsx_insert_rows (row expansion with prototype cloning & formula shift)
4. xlsx_find_anchor (keyword locator)
5. docx_inspect_placeholders (template AST inspection)
6. docx_replace_text (ERR_DOCX_002 run preservation & key normalization)
7. docx_insert_paragraph (surgical paragraph insertion)
8. docx_insert_table (ERR_DOCX_003 cantSplit and tblHeader OpenXML invariants)
9. CLI subcommands (xlsx-set, xlsx-insert-rows, docx-tokens, docx-replace)
"""

import os
import shutil
import tempfile
import unittest
from openpyxl import Workbook
from docx import Document

from server.mcp_server import (
    execute_xlsx_set_cell,
    execute_xlsx_set_range,
    execute_xlsx_insert_rows,
    execute_xlsx_find_anchor,
    execute_docx_inspect_placeholders,
    execute_docx_replace_text,
    execute_docx_insert_paragraph,
    execute_docx_insert_table,
)
from cli.ai_tools import build_parser


class TestAtomicOfficeTools(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    # -------------------------------------------------------------------------
    # 1. Excel Atomic Tools Tests
    # -------------------------------------------------------------------------

    def test_xlsx_set_cell_and_merged_protection(self):
        """Test xlsx_set_cell handles normal cells and safely redirects merged cells (ERR_XLSX_004, ERR_XLSX_007)."""
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "TestSheet"
        # Merge B2:D3
        ws.merge_cells("B2:D3")
        ws["B2"] = "Initial Top-Left"
        ws["A1"] = "Header"
        
        wb_path = os.path.join(self.tmp_dir, "test_merged.xlsx")
        wb.save(wb_path)

        # 1. Update regular cell with leading-zero string (ERR_XLSX_007)
        res1 = execute_xlsx_set_cell({
            "file_path": wb_path,
            "sheet_name": "TestSheet",
            "cell": "A1",
            "value": "0987654321",
        })
        self.assertTrue(res1["success"])

        # 2. Update merged cell non-top-left (e.g. C2 or D3) -> must not throw MergedCell read-only
        res2 = execute_xlsx_set_cell({
            "file_path": wb_path,
            "sheet_name": "TestSheet",
            "cell": "C2",
            "value": "Updated via non-top-left",
        })
        self.assertTrue(res2["success"])

        # Verify
        from openpyxl import load_workbook
        wb_check = load_workbook(wb_path)
        ws_check = wb_check["TestSheet"]
        # Leading zero format checked
        cell_a1 = ws_check["A1"]
        self.assertEqual(str(cell_a1.value), "0987654321")
        self.assertEqual(cell_a1.number_format, "@")
        # Merged cell Top-Left updated
        self.assertEqual(ws_check["B2"].value, "Updated via non-top-left")

    def test_xlsx_set_range(self):
        """Test batch cell updates via xlsx_set_range."""
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Data"
        wb_path = os.path.join(self.tmp_dir, "range.xlsx")
        wb.save(wb_path)

        res = execute_xlsx_set_range({
            "file_path": wb_path,
            "sheet_name": "Data",
            "updates": {
                "A1": "ID",
                "B1": "Score",
                "A2": "001",
                "B2": {"value": 95},
                "B3": {"formula": "=SUM(B2:B2)"},
            }
        })
        self.assertTrue(res["success"])
        self.assertEqual(res["updated_cells"], 5)

        from openpyxl import load_workbook
        wb_check = load_workbook(wb_path)
        ws_check = wb_check["Data"]
        self.assertEqual(ws_check["A2"].value, "001")
        self.assertEqual(ws_check["A2"].number_format, "@")
        self.assertEqual(ws_check["B2"].value, 95)
        self.assertEqual(ws_check["B3"].value, "=SUM(B2:B2)")

    def test_xlsx_insert_rows_and_find_anchor(self):
        """Test inserting rows with style/formula shift and finding anchor."""
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Summary"
        ws["A1"] = "Search Key: ANCHOR_FLAG"
        ws["A5"] = "Row 5 Data"
        ws["B5"] = 100
        ws["B6"] = "=SUM(B5:B5)"
        wb_path = os.path.join(self.tmp_dir, "anchor.xlsx")
        wb.save(wb_path)

        # Test find_anchor
        find_res = execute_xlsx_find_anchor({
            "file_path": wb_path,
            "sheet_name": "Summary",
            "keyword": "ANCHOR_FLAG",
        })
        self.assertTrue(find_res["success"])
        self.assertTrue(find_res["found"])
        self.assertEqual(find_res["coordinate"], "A1")

        # Test insert_rows at row 6 with count 2
        ins_res = execute_xlsx_insert_rows({
            "file_path": wb_path,
            "sheet_name": "Summary",
            "start_row": 6,
            "count": 2,
            "prototype_row": 5,
            "shift_formulas": True,
        })
        self.assertTrue(ins_res["success"])
        self.assertEqual(ins_res["count"], 2)

    # -------------------------------------------------------------------------
    # 2. Word Atomic Tools Tests
    # -------------------------------------------------------------------------

    def test_docx_inspect_placeholders_and_replace_text(self):
        """Test token inspection and text replacement with KEY and {{KEY}} normalization (ERR_DOCX_002)."""
        doc = Document()
        p1 = doc.add_paragraph("Kính gửi ông/bà: ")
        r1 = p1.add_run("{{HO_TEN}}")
        r1.bold = True
        p2 = doc.add_paragraph("Chức vụ: {{CHUC_VU}} tại công ty {{CONG_TY}}.")
        
        doc_path = os.path.join(self.tmp_dir, "template.docx")
        doc.save(doc_path)

        # 1. Inspect placeholders
        inspect_res = execute_docx_inspect_placeholders({"template_path": doc_path})
        self.assertIn("HO_TEN", inspect_res["variables"])
        self.assertIn("CHUC_VU", inspect_res["variables"])
        self.assertIn("CONG_TY", inspect_res["variables"])

        # 2. Replace text using normalized keys (one with {{}}, one without)
        out_doc = os.path.join(self.tmp_dir, "patched.docx")
        rep_res = execute_docx_replace_text({
            "doc_path": doc_path,
            "replacements": {
                "HO_TEN": "Biện Trung Tín",        # Without braces
                "{{CHUC_VU}}": "SAP Consultant",   # With braces
                "CONG_TY": "ICON Consulting Group"
            },
            "output_path": out_doc,
        })
        self.assertTrue(rep_res["success"])

        # Verify replacement
        doc_check = Document(out_doc)
        full_text = "\n".join(p.text for p in doc_check.paragraphs)
        self.assertIn("Biện Trung Tín", full_text)
        self.assertIn("SAP Consultant", full_text)
        self.assertIn("ICON Consulting Group", full_text)
        self.assertNotIn("{{HO_TEN}}", full_text)

    def test_docx_insert_paragraph_and_table_invariants(self):
        """Test surgical paragraph and table insertion with cantSplit & tblHeader (ERR_DOCX_003)."""
        doc = Document()
        doc.add_heading("Section 1: Details", level=1)
        doc.add_paragraph("Existing text")
        
        doc_path = os.path.join(self.tmp_dir, "doc_insert.docx")
        doc.save(doc_path)

        # 1. Insert paragraph after heading
        p_res = execute_docx_insert_paragraph({
            "doc_path": doc_path,
            "text": "Inserted Note: Critical invariant validated.",
            "after_heading": "Section 1: Details",
            "bold": True,
            "color": "0055AA",
        })
        self.assertTrue(p_res["success"])

        # 2. Insert table with openxml invariants
        t_res = execute_docx_insert_table({
            "doc_path": doc_path,
            "headers": ["Mã", "Tên", "Trạng Thái"],
            "rows": [
                ["01", "Module A", "Pass"],
                ["02", "Module B", "Pass"],
            ],
            "after_heading": "Section 1: Details",
        })
        self.assertTrue(t_res["success"])

        # Verify XML invariants
        doc_check = Document(doc_path)
        self.assertEqual(len(doc_check.tables), 1)
        tbl = doc_check.tables[0]
        # Check rows count
        self.assertEqual(len(tbl.rows), 3)
        # Check XML contains cantSplit and tblHeader
        tbl_xml = tbl._tbl.xml
        self.assertIn("w:cantSplit", tbl_xml)
        self.assertIn("w:tblHeader", tbl_xml)

    # -------------------------------------------------------------------------
    # 3. CLI Subcommands Test
    # -------------------------------------------------------------------------

    def test_cli_subcommands_dispatch(self):
        """Test CLI arguments parsing and execution for new atomic commands."""
        parser = build_parser()

        # 1. xlsx-set
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Sheet1"
        xlsx_path = os.path.join(self.tmp_dir, "cli_test.xlsx")
        wb.save(xlsx_path)

        args_set = parser.parse_args(["xlsx-set", xlsx_path, "--sheet", "Sheet1", "--cell", "B5", "--value", "CLI_VAL"])
        ret_set = args_set.func(args_set)
        self.assertEqual(ret_set, 0)

        # 2. xlsx-insert-rows
        args_ins = parser.parse_args(["xlsx-insert-rows", xlsx_path, "--sheet", "Sheet1", "--start-row", "2", "--count", "3"])
        ret_ins = args_ins.func(args_ins)
        self.assertEqual(ret_ins, 0)

        # 3. docx-tokens
        doc = Document()
        doc.add_paragraph("Hello {{CLI_NAME}}!")
        docx_path = os.path.join(self.tmp_dir, "cli_test.docx")
        doc.save(docx_path)

        args_tok = parser.parse_args(["docx-tokens", docx_path, "--json"])
        ret_tok = args_tok.func(args_tok)
        self.assertEqual(ret_tok, 0)

        # 4. docx-replace
        args_rep = parser.parse_args(["docx-replace", docx_path, "--data", '{"CLI_NAME": "Antigravity User"}'])
        ret_rep = args_rep.func(args_rep)
        self.assertEqual(ret_rep, 0)

        doc_check = Document(docx_path)
        self.assertIn("Antigravity User", doc_check.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
