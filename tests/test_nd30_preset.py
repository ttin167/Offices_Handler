"""
test_nd30_preset.py — Unit Tests for Decree 30/2020/ND-CP Presets & Vietnamese Typography.

Tests:
1. Preset discovery and configuration loading (ND30, Academic Thesis, Corporate Report).
2. Document geometry and typography application (A4, 20/30/20mm margins, Times New Roman 13pt).
3. Decree 30 administrative header generation with OpenXML invariants (cantSplit, tblHeader, ERR_DOCX_001).
4. VietnameseTypographySanitizer (Unicode NFC, punctuation spacing, bracket trimming, En-dash ranges).
5. Document-wide typography sanitization across paragraphs, tables, and headers.
"""

import os
import sys
import unittest
import unicodedata
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from core.docx.presets import (
    list_available_presets,
    load_preset,
    apply_preset_to_document,
    build_administrative_header,
    BUILTIN_ND30_PRESET,
)
from core.docx.advanced_engine import (
    VietnameseTypographySanitizer,
    sanitize_vietnamese_typography,
)


class TestND30PresetAndTypography(unittest.TestCase):
    def setUp(self):
        self.scratch_dir = os.path.join(BASE_DIR, "scratch", "test_nd30")
        os.makedirs(self.scratch_dir, exist_ok=True)

    def test_preset_discovery_and_loading(self):
        """Test listing available presets and loading configuration files."""
        presets = list_available_presets()
        self.assertIn("nd30", presets)
        self.assertIn("academic_thesis", presets)

        nd30_cfg = load_preset("nd30")
        self.assertEqual(nd30_cfg["geometry"]["page_size"], "A4")
        self.assertEqual(nd30_cfg["geometry"]["margin"]["left"], "30mm")
        self.assertEqual(nd30_cfg["geometry"]["margin"]["top"], "20mm")
        self.assertIn("Times New Roman", nd30_cfg["typography"]["font_family"])

        # Fallback test for unknown preset
        fallback_cfg = load_preset("non_existent_preset_xyz")
        self.assertEqual(fallback_cfg["geometry"]["page_size"], "A4")

    def test_apply_preset_to_document(self):
        """Test applying ND30 preset modifies margins, page size, and Normal style."""
        doc = Document()
        res = apply_preset_to_document(doc, "nd30")
        self.assertTrue(res["success"])

        for sec in doc.sections:
            self.assertAlmostEqual(sec.top_margin.mm, 20.0, places=1)
            self.assertAlmostEqual(sec.bottom_margin.mm, 20.0, places=1)
            self.assertAlmostEqual(sec.left_margin.mm, 30.0, places=1)
            self.assertAlmostEqual(sec.right_margin.mm, 20.0, places=1)
            self.assertAlmostEqual(sec.page_width.mm, 210.0, places=1)
            self.assertAlmostEqual(sec.page_height.mm, 297.0, places=1)

        normal_style = doc.styles["Normal"]
        self.assertEqual(normal_style.font.name, "Times New Roman")
        self.assertAlmostEqual(normal_style.font.size.pt, 13.0, places=1)
        self.assertAlmostEqual(normal_style.paragraph_format.line_spacing, 1.3, places=1)

    def test_build_administrative_header(self):
        """Test administrative header table structure and OpenXML invariants."""
        doc = Document()
        table = build_administrative_header(
            doc,
            agency_parent="BỘ KHOA HỌC VÀ CÔNG NGHỆ",
            agency_name="CỤC BẢN QUYỀN",
            document_no="Số: 42/BPS-KHCN",
            location_date="Hà Nội, ngày 18 tháng 10 năm 2026",
            republic_title="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            motto="Độc lập - Tự do - Hạnh phúc",
            insert_at_top=True,
        )

        self.assertEqual(len(table.rows), 1)
        self.assertEqual(len(table.columns), 2)

        # Invariant check: cantSplit and tblHeader in row trPr
        row = table.rows[0]
        trPr = row._tr.find(qn("w:trPr"))
        self.assertIsNotNone(trPr)
        self.assertIsNotNone(trPr.find(qn("w:cantSplit")))
        self.assertIsNotNone(trPr.find(qn("w:tblHeader")))

        # Invariant check: ERR_DOCX_001 (all table cells must end with <w:p>)
        for cell in row.cells:
            last_elem = cell._tc[-1]
            self.assertEqual(last_elem.tag, qn("w:p"))

        # Verify agency text and motto present
        cell_0_text = "\n".join(p.text for p in row.cells[0].paragraphs)
        cell_1_text = "\n".join(p.text for p in row.cells[1].paragraphs)

        self.assertIn("BỘ KHOA HỌC VÀ CÔNG NGHỆ", cell_0_text)
        self.assertIn("CỤC BẢN QUYỀN", cell_0_text)
        self.assertIn("Số: 42/BPS-KHCN", cell_0_text)

        self.assertIn("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", cell_1_text)
        self.assertIn("Độc lập - Tự do - Hạnh phúc", cell_1_text)
        self.assertIn("Hà Nội, ngày 18 tháng 10 năm 2026", cell_1_text)

        # Save and verify valid DOCX output
        out_path = os.path.join(self.scratch_dir, "test_admin_doc.docx")
        doc.save(out_path)
        self.assertTrue(os.path.isfile(out_path))

    def test_vietnamese_typography_sanitizer_rules(self):
        """Test atomic string sanitization rules in VietnameseTypographySanitizer."""
        # 1. Unicode NFC Normalization (composite -> precomposed)
        decomposed = "Viê\u0323t Nam"  # 'ệ' decomposed
        normalized = VietnameseTypographySanitizer.sanitize_string(decomposed)
        self.assertEqual(unicodedata.normalize("NFC", normalized), normalized)
        self.assertEqual(normalized, "Việt Nam")

        # 2. Punctuation spacing (no space before, space after)
        raw_punct = "Xin chào ,thế giới !Đây là thử nghiệm ;về văn bản ."
        cleaned_punct = VietnameseTypographySanitizer.sanitize_string(raw_punct)
        self.assertEqual(cleaned_punct, "Xin chào, thế giới! Đây là thử nghiệm; về văn bản.")

        # 3. Bracket and quote snug spacing
        raw_brackets = "Thông tư số ( 30/2020/TT-BTTTT ) và “ trích dẫn mẫu ”."
        cleaned_brackets = VietnameseTypographySanitizer.sanitize_string(raw_brackets)
        self.assertEqual(cleaned_brackets, "Thông tư số (30/2020/TT-BTTTT) và “trích dẫn mẫu”.")

        # 4. En-dash normalization for numeric/year ranges
        raw_ranges = "Kế hoạch 2021-2025 và giai đoạn 2026 - 2030."
        cleaned_ranges = VietnameseTypographySanitizer.sanitize_string(raw_ranges)
        self.assertIn("2021 – 2025", cleaned_ranges)
        self.assertIn("2026 – 2030", cleaned_ranges)

        # 5. Multiple space collapsing
        raw_spaces = "Cộng hòa    xã hội   chủ nghĩa  Việt   Nam"
        cleaned_spaces = VietnameseTypographySanitizer.sanitize_string(raw_spaces)
        self.assertEqual(cleaned_spaces, "Cộng hòa xã hội chủ nghĩa Việt Nam")

    def test_sanitize_vietnamese_typography_in_document(self):
        """Test document-wide sanitization modifying paragraphs and table cells."""
        doc = Document()
        p1 = doc.add_paragraph("Kế hoạch năm 2025-2026 ,được ban hành theo Nghị định .")
        table = doc.add_table(rows=1, cols=1)
        p2 = table.rows[0].cells[0].paragraphs[0]
        p2.text = "Giai đoạn 2020-2024 ;đã hoàn thành ( 100% ) ."

        mod_count = sanitize_vietnamese_typography(doc)
        self.assertGreater(mod_count, 0)

        self.assertIn("2025 – 2026, được ban hành theo Nghị định.", p1.text)
        self.assertIn("2020 – 2024; đã hoàn thành (100%).", p2.text)


if __name__ == "__main__":
    unittest.main()
