"""
test_mail_merge.py — Unit Tests for Enterprise Batch Mail Merge Studio.

Verifies:
1. inspect_template_tokens: Full token, conditional, and block discovery across body, tables, headers.
2. evaluate_condition: Safe evaluation of equality, numerical ranges, membership, and truthiness.
3. resolve_conditional_text: In-place {{IF:...}}...{{ELSE}}...{{ENDIF}} substitution.
4. read_data_source: Tabular reading from Excel/CSV/JSON with leading-zero preservation (ERR_XLSX_007).
5. execute_mail_merge: End-to-end batch generation, dynamic filename pattern, and ZIP bundle packaging.
"""

import datetime
import os
import sys
import unittest
import zipfile
from docx import Document
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from docx_mail_merge import (
    inspect_template_tokens,
    evaluate_condition,
    resolve_conditional_text,
    read_data_source,
    execute_mail_merge,
)


class TestMailMergeStudio(unittest.TestCase):
    def setUp(self):
        self.scratch_dir = os.path.join(BASE_DIR, "scratch", "test_mail_merge")
        os.makedirs(self.scratch_dir, exist_ok=True)

    def test_evaluate_condition_and_resolve_text(self):
        """Test declarative evaluation of condition expressions and text substitution."""
        rec1 = {"Gender": "Nam", "Age": 25, "Score": 9.2, "IsVip": True, "Role": "Admin"}
        rec2 = {"Gender": "Nu", "Age": 16, "Score": 6.5, "IsVip": False, "Role": "Guest"}

        # 1. Condition evaluation
        self.assertTrue(evaluate_condition("Gender == 'Nam'", rec1))
        self.assertFalse(evaluate_condition("Gender == 'Nam'", rec2))
        self.assertTrue(evaluate_condition("Age >= 18", rec1))
        self.assertFalse(evaluate_condition("Age >= 18", rec2))
        self.assertTrue(evaluate_condition("Score > 8.0", rec1))
        self.assertFalse(evaluate_condition("Score > 8.0", rec2))
        self.assertTrue(evaluate_condition("IsVip", rec1))
        self.assertFalse(evaluate_condition("IsVip", rec2))
        self.assertTrue(evaluate_condition("Role in ['Admin', 'Lead']", rec1))
        self.assertFalse(evaluate_condition("Role in ['Admin', 'Lead']", rec2))

        # 2. Text substitution with {{IF:...}}...{{ELSE}}...{{ENDIF}}
        tmpl = "Kính gửi {{IF:Gender == 'Nam'}}Ông{{ELSE}}Bà{{ENDIF}} {{HoTen}}, trạng thái: {{IF:Age >= 18}}Người lớn{{ELSE}}Vị thành niên{{ENDIF}}."
        res1 = resolve_conditional_text(tmpl, {"Gender": "Nam", "Age": 25, "HoTen": "Nam"})
        self.assertEqual(res1, "Kính gửi Ông {{HoTen}}, trạng thái: Người lớn.")

        res2 = resolve_conditional_text(tmpl, {"Gender": "Nu", "Age": 15, "HoTen": "Lan"})
        self.assertEqual(res2, "Kính gửi Bà {{HoTen}}, trạng thái: Vị thành niên.")

    def test_inspect_template_tokens(self):
        """Test discovering all tokens across body, tables, and headers/footers."""
        template_file = os.path.join(self.scratch_dir, "sample_template.docx")
        doc = Document()

        # Body
        doc.add_paragraph("Hợp đồng lao động số {{SoHD}} giữa {{TenCongTy}} và {{HoTen}}.")
        doc.add_paragraph("Điều kiện: {{IF:Luong > 15000000}}Áp dụng thuế TNCN{{ENDIF}}.")
        doc.add_paragraph("{{BLOCK:BangChiTiet}}")

        # Table
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].paragraphs[0].text = "Chức vụ: {{ChucVu}}"
        t.rows[0].cells[1].paragraphs[0].text = "Mã NV: {{MaNV}}"

        # Header
        doc.sections[0].header.paragraphs[0].text = "Mẫu chuẩn 2026 — {{PhongBan}}"

        doc.save(template_file)

        info = inspect_template_tokens(template_file)
        self.assertIn("SoHD", info["variables"])
        self.assertIn("TenCongTy", info["variables"])
        self.assertIn("HoTen", info["variables"])
        self.assertIn("ChucVu", info["variables"])
        self.assertIn("MaNV", info["variables"])
        self.assertIn("PhongBan", info["variables"])
        self.assertIn("Luong > 15000000", info["conditionals"])
        self.assertIn("BangChiTiet", info["blocks"])

    def test_read_data_source_excel_leading_zero(self):
        """Test reading Excel data source preserving leading zeros and dates (ERR_XLSX_007)."""
        xlsx_path = os.path.join(self.scratch_dir, "employees.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "NhanVien"

        ws.append(["Mã NV", "Họ và Tên", "Số ĐT", "Ngày Vào", "Lương"])
        # Format Mã NV as string with leading zeroes
        ws.cell(row=2, column=1, value="00128").number_format = "@"
        ws.cell(row=2, column=2, value="Nguyen Van An")
        ws.cell(row=2, column=3, value="0912345678").number_format = "@"
        ws.cell(row=2, column=4, value=datetime.date(2026, 5, 1))
        ws.cell(row=2, column=5, value=25000000)

        wb.save(xlsx_path)

        records = read_data_source(xlsx_path)
        self.assertEqual(len(records), 1)
        r = records[0]
        self.assertEqual(r["Mã NV"], "00128")
        self.assertEqual(r["Số ĐT"], "0912345678")
        self.assertEqual(r["Họ và Tên"], "Nguyen Van An")
        self.assertEqual(r["Ngày Vào"], "01/05/2026")

    def test_execute_mail_merge_full_batch_and_zip(self):
        """Test end-to-end batch generation and zip bundle creation."""
        template_file = os.path.join(self.scratch_dir, "template_merge.docx")
        doc = Document()
        doc.add_paragraph("Kính gửi {{IF:GioiTinh == 'Nam'}}Ông{{ELSE}}Bà{{ENDIF}} {{HoTen}},")
        doc.add_paragraph("Mã nhân viên: {{MaNV}}")
        doc.add_paragraph("Chức vụ bổ nhiệm: {{ChucVu}}")
        doc.add_paragraph("Quyết định có hiệu lực từ ngày 19/10/2026.")
        doc.save(template_file)

        data = [
            {"MaNV": "NV001", "HoTen": "Tran Van Nam", "GioiTinh": "Nam", "ChucVu": "Giam Doc Kỹ Thuat"},
            {"MaNV": "NV002", "HoTen": "Le Thi Mai", "GioiTinh": "Nu", "ChucVu": "Truong Phong Nhan Su"},
            {"MaNV": "NV003", "HoTen": "Pham Quoc Bao", "GioiTinh": "Nam", "ChucVu": "Kien Truc Su He Thong"},
        ]

        out_dir = os.path.join(self.scratch_dir, "merged_output")
        res = execute_mail_merge(
            template_path=template_file,
            data_source=data,
            output_dir=out_dir,
            filename_pattern="QuyetDinh_{{MaNV}}_{{HoTen}}.docx",
            create_zip=True,
            zip_name="quyet_dinh_bundle.zip",
        )

        self.assertTrue(res["success"])
        self.assertEqual(res["total_records"], 3)
        self.assertEqual(len(res["generated_files"]), 3)
        self.assertIsNotNone(res["zip_path"])
        self.assertTrue(os.path.isfile(res["zip_path"]))

        # Verify generated file 1
        doc1 = Document(res["generated_files"][0])
        full_text1 = "\n".join(p.text for p in doc1.paragraphs)
        self.assertIn("Kính gửi Ông Tran Van Nam,", full_text1)
        self.assertIn("Mã nhân viên: NV001", full_text1)
        self.assertIn("Chức vụ bổ nhiệm: Giam Doc Kỹ Thuat", full_text1)

        # Verify generated file 2
        doc2 = Document(res["generated_files"][1])
        full_text2 = "\n".join(p.text for p in doc2.paragraphs)
        self.assertIn("Kính gửi Bà Le Thi Mai,", full_text2)
        self.assertIn("Mã nhân viên: NV002", full_text2)

        # Verify ZIP bundle
        with zipfile.ZipFile(res["zip_path"], "r") as zf:
            namelist = zf.namelist()
            self.assertEqual(len(namelist), 3)
            self.assertTrue(any("NV001" in n for n in namelist))
            self.assertTrue(any("NV002" in n for n in namelist))


if __name__ == "__main__":
    unittest.main()
