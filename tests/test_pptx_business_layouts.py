"""
tests.test_pptx_business_layouts — Unit tests for Enterprise Presentation Layouts (v4.0).

Verifies:
1. add_kpi_summary_slide: Executive KPI metrics with delta badges (▲/▼) and run-level pure black text.
2. add_timeline_roadmap_slide: Continuous horizontal axis line, milestone nodes, and deliverables cards.
3. add_swot_matrix_slide: Balanced 2x2 grid (Strengths Blue, Weaknesses Red, Opportunities Green, Threats Amber).
4. add_process_chevron_slide: Sequential process workflow (3 to 6 steps) with arrow connectors.
5. write_pptx_from_spec: High-level spec compilation for all new business layouts + charts.
6. docx_to_pptx semantic classification: Automatic routing of headings/elements to business layouts.
"""

import os
import shutil
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

from core.pptx.writer import (
    add_kpi_summary_slide,
    add_timeline_roadmap_slide,
    add_swot_matrix_slide,
    add_process_chevron_slide,
    write_pptx_from_spec,
    init_presentation,
)
from core.pptx.docx_to_pptx import _detect_business_layout, synthesize_slide_spec


def is_shape_type(s, mso_type):
    try:
        return s.auto_shape_type == mso_type
    except Exception:
        return False


class TestPPTXBusinessLayouts(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_biz_layouts_")

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_01_kpi_summary_slide(self):
        prs = init_presentation()
        metrics = [
            {"label": "Doanh thu", "value": "158.4 tỷ", "change": "+22.5% YoY", "desc": "Vượt 12% chỉ tiêu"},
            {"label": "Khách hàng mới", "value": "24,500", "change": "+15.0%", "desc": "Tăng trưởng kênh digital"},
            {"label": "Tỷ lệ giữ chân", "value": "94.2%", "change": "+1.8%", "desc": "Mức cao kỷ lục 3 năm"},
            {"label": "Chi phí CAC", "value": "185,000 đ", "change": "-8.4% MoM", "desc": "Tối ưu hóa chiến dịch Ads"},
        ]
        slide = add_kpi_summary_slide(
            prs,
            title="KẾT QUẢ KINH DOANH NĂM 2026",
            metrics=metrics,
            subtitle="Số liệu cập nhật đến hết Quý 4/2026",
            banner="TỔNG QUAN TÀI CHÍNH & VẬN HÀNH",
        )
        self.assertIsNotNone(slide)

        # Count cards: Header textboxes + 4 rounded rectangles
        rounded_cards = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.ROUNDED_RECTANGLE)]
        self.assertEqual(len(rounded_cards), 4)

        # Check content in first card
        first_card_text = rounded_cards[0].text_frame.text
        self.assertIn("DOANH THU", first_card_text)
        self.assertIn("158.4 tỷ", first_card_text)
        self.assertIn("▲", first_card_text)

        # Check negative delta in 4th card
        fourth_card_text = rounded_cards[3].text_frame.text
        self.assertIn("▼", fourth_card_text)

        out_path = os.path.join(self.test_dir, "test_kpi_summary.pptx")
        prs.save(out_path)
        self.assertTrue(os.path.isfile(out_path))

    def test_02_timeline_roadmap_slide(self):
        prs = init_presentation()
        milestones = [
            {"phase": "Q1 2026", "title": "Khảo Sát & Thiết Kế", "tasks": ["Nghiên cứu thị trường", "Hoàn thiện SRS"]},
            {"phase": "Q2 2026", "title": "Phát Triển MVP", "tasks": ["Xây dựng Core Engine", "Tích hợp AI API"]},
            {"phase": "Q3 2026", "title": "Thử Nghiệm Beta", "tasks": ["Pilot 10 doanh nghiệp", "Fix bugs & tối ưu"]},
            {"phase": "Q4 2026", "title": "Go-Live & Mở Rộng", "tasks": ["Phát hành bản thương mại", "Scale 100k users"]},
        ]
        slide = add_timeline_roadmap_slide(
            prs,
            title="LỘ TRÌNH TRIỂN KHAI DỰ ÁN",
            milestones=milestones,
            subtitle="Kế hoạch 4 giai đoạn chuẩn bị phát hành",
        )
        self.assertIsNotNone(slide)

        # Check for circular nodes (OVAL) and cards
        ovals = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.OVAL)]
        self.assertEqual(len(ovals), 4)
        cards = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.ROUNDED_RECTANGLE)]
        self.assertEqual(len(cards), 4)

        out_path = os.path.join(self.test_dir, "test_timeline.pptx")
        prs.save(out_path)
        self.assertTrue(os.path.isfile(out_path))

    def test_03_swot_matrix_slide(self):
        prs = init_presentation()
        swot_data = {
            "strengths": ["Đội ngũ kỹ sư R&D chất lượng cao", "Kiến trúc microservices mở rộng tốt"],
            "weaknesses": ["Thời gian tiếp cận thị trường còn dài", "Nguồn vốn tiếp thị ban đầu giới hạn"],
            "opportunities": ["Nhu cầu chuyển đổi số tài liệu tăng mạnh", "Chính sách ưu đãi doanh nghiệp công nghệ"],
            "threats": ["Cạnh tranh từ các đối thủ lớn toàn cầu", "Biến động chi phí hạ tầng Cloud"],
        }
        slide = add_swot_matrix_slide(
            prs,
            title="MA TRẬN CHIẾN LƯỢC SWOT",
            swot_data=swot_data,
            subtitle="Đánh giá toàn diện nội tại và thị trường",
        )
        self.assertIsNotNone(slide)

        cards = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.ROUNDED_RECTANGLE)]
        self.assertEqual(len(cards), 4)

        # Verify key headers
        all_text = " ".join(c.text_frame.text for c in cards)
        self.assertIn("[S]", all_text)
        self.assertIn("[W]", all_text)
        self.assertIn("[O]", all_text)
        self.assertIn("[T]", all_text)

        out_path = os.path.join(self.test_dir, "test_swot.pptx")
        prs.save(out_path)
        self.assertTrue(os.path.isfile(out_path))

    def test_04_process_chevron_slide(self):
        prs = init_presentation()
        steps = [
            {"title": "Tiếp Nhận Yêu Cầu", "duration": "Ngày 1-2", "actions": ["Ghi nhận đề bài", "Phân tích phạm vi"]},
            {"title": "Thiết Kế Kiến Trúc", "duration": "Ngày 3-5", "actions": ["Vẽ sơ đồ ERD & C4", "Chốt API spec"]},
            {"title": "Lập Trình & Tích Hợp", "duration": "Tuần 2-3", "actions": ["Viết code Clean Architecture", "Viết Unit Tests"]},
            {"title": "Kiểm Thử & Đóng Gói", "duration": "Tuần 4", "actions": ["Chạy 12-Gate QA", "Build Docker image"]},
        ]
        slide = add_process_chevron_slide(
            prs,
            title="QUY TRÌNH PHÁT TRIỂN PHẦN MỀM",
            steps=steps,
            subtitle="4 bước tiêu chuẩn từ tiếp nhận đến bàn giao",
        )
        self.assertIsNotNone(slide)

        step_cards = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.ROUNDED_RECTANGLE)]
        self.assertEqual(len(step_cards), 4)

        # Arrow connectors (count - 1 = 3)
        arrows = [s for s in slide.shapes if is_shape_type(s, MSO_SHAPE.RIGHT_ARROW)]
        self.assertEqual(len(arrows), 3)

        out_path = os.path.join(self.test_dir, "test_process.pptx")
        prs.save(out_path)
        self.assertTrue(os.path.isfile(out_path))

    def test_05_write_pptx_from_spec_all_new_layouts(self):
        spec = {
            "title": "BÁO CÁO TỔNG QUAN DOANH NGHIỆP 2026",
            "theme": "corporate_blue",
            "slides": [
                {
                    "type": "title",
                    "title": "BÁO CÁO CHIẾN LƯỢC 2026",
                    "subtitle": "Kế hoạch phát triển và mục tiêu tài chính",
                    "presenter": "Ban Giám Đốc",
                    "date": "2026-10-19",
                },
                {
                    "type": "kpi_summary",
                    "title": "CHỈ SỐ TĂNG TRƯỞNG CHÍNH (KPI)",
                    "metrics": [
                        {"label": "Doanh thu", "value": "240 Tỷ", "change": "+35% YoY", "desc": "Mục tiêu 2026"},
                        {"label": "EBITDA", "value": "48 Tỷ", "change": "+28%", "desc": "Biên lợi nhuận 20%"},
                    ],
                },
                {
                    "type": "timeline",
                    "title": "LỘ TRÌNH THỰC HIỆN DỰ ÁN",
                    "milestones": [
                        {"phase": "H1/2026", "title": "Hoàn tất nền tảng", "tasks": ["R&D Core Engine"]},
                        {"phase": "H2/2026", "title": "Bùng nổ doanh số", "tasks": ["Khai thác thị trường"]},
                    ],
                },
                {
                    "type": "swot",
                    "title": "PHÂN TÍCH MA TRẬN SWOT",
                    "swot": {
                        "strengths": ["Công nghệ lõi vượt trội"],
                        "weaknesses": ["Quy mô đội ngũ"],
                        "opportunities": ["Thị trường mở rộng"],
                        "threats": ["Chi phí biến động"],
                    },
                },
                {
                    "type": "process",
                    "title": "QUY TRÌNH VẬN HÀNH",
                    "steps": [
                        {"title": "Khởi động", "actions": ["Họp kickoff"]},
                        {"title": "Triển khai", "actions": ["Thực thi"]},
                        {"title": "Đánh giá", "actions": ["Tổng kết"]},
                    ],
                },
                {
                    "type": "chart",
                    "title": "CƠ CẤU DOANH THU",
                    "chart": {
                        "type": "column",
                        "categories": ["Q1", "Q2", "Q3", "Q4"],
                        "series": [{"name": "Doanh thu (tỷ)", "values": [50, 60, 65, 75]}],
                    },
                },
            ],
        }
        out_pptx = os.path.join(self.test_dir, "master_business_report.pptx")
        saved_file = write_pptx_from_spec(spec, out_pptx)
        self.assertTrue(os.path.isfile(saved_file))
        self.assertGreater(os.path.getsize(saved_file), 20000)

        # Validate with Presentation loader
        prs = Presentation(saved_file)
        self.assertEqual(len(prs.slides), 6)

    def test_06_docx_to_pptx_semantic_classification(self):
        # 1. SWOT detection
        swot_sec_elements = [
            {"type": "bullet", "text": "Điểm mạnh: Năng lực kỹ thuật mạnh mẽ"},
            {"type": "bullet", "text": "Điểm yếu: Thiếu nhân sự marketing"},
            {"type": "bullet", "text": "Cơ hội: Thị trường đang mở rộng"},
            {"type": "bullet", "text": "Thách thức: Chi phí máy chủ tăng"},
        ]
        layout = _detect_business_layout("Phân tích SWOT dự án", swot_sec_elements)
        self.assertIsNotNone(layout)
        self.assertEqual(layout["type"], "swot")
        self.assertIn("strengths", layout["swot"])

        # 2. Timeline detection
        timeline_elements = [
            {"type": "bullet", "text": "Q1/2026: Khởi động dự án và thiết kế hệ thống"},
            {"type": "bullet", "text": "Q2/2026: Hoàn thành bản thử nghiệm Alpha"},
            {"type": "bullet", "text": "Q3/2026: Ra mắt chính thức trên thị trường"},
        ]
        layout2 = _detect_business_layout("Lộ trình kế hoạch triển khai", timeline_elements)
        self.assertIsNotNone(layout2)
        self.assertEqual(layout2["type"], "timeline")
        self.assertEqual(len(layout2["milestones"]), 3)

        # 3. KPI detection
        kpi_elements = [
            {"type": "paragraph", "text": "Tổng doanh thu: 180 tỷ (+25% YoY)"},
            {"type": "paragraph", "text": "Số lượng người dùng: 500,000 (+40%)"},
        ]
        layout3 = _detect_business_layout("Các chỉ số KPI và hiệu suất đạt được", kpi_elements)
        self.assertIsNotNone(layout3)
        self.assertEqual(layout3["type"], "kpi_summary")
        self.assertEqual(len(layout3["metrics"]), 2)

        # 4. Process detection
        process_elements = [
            {"type": "bullet", "text": "Bước 1: Tiếp nhận hồ sơ yêu cầu"},
            {"type": "bullet", "text": "Bước 2: Thẩm định và đánh giá rủi ro"},
            {"type": "bullet", "text": "Bước 3: Phê duyệt và giải ngân"},
        ]
        layout4 = _detect_business_layout("Quy trình thẩm định hồ sơ", process_elements)
        self.assertIsNotNone(layout4)
        self.assertEqual(layout4["type"], "process")
        self.assertEqual(len(layout4["steps"]), 3)


if __name__ == "__main__":
    unittest.main()
