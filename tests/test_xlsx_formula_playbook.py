# -*- coding: utf-8 -*-
"""
tests.test_xlsx_formula_playbook — Test suite for Excel Formula Playbook & Cheat Sheet System.
"""

from __future__ import annotations

import argparse
import io
import os
import sys
import tempfile
import unittest

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

from core.xlsx.formula_playbook import (
    FORMULA_CATALOG,
    get_formula_catalog,
    generate_formula_playbook_xlsx,
)
from server.mcp_server import execute_xlsx_formula_helper
from cli.ai_tools import cmd_xlsx_formulas


class TestExcelFormulaPlaybook(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_catalog_integrity_and_required_fields(self):
        """Ensure all 32 formulas contain complete 4-dimension attributes without placeholders."""
        self.assertGreaterEqual(len(FORMULA_CATALOG), 32)
        required_keys = {
            "name",
            "category",
            "category_vi",
            "template",
            "example",
            "when_to_use",
            "why_choose",
        }

        categories = set()
        for idx, item in enumerate(FORMULA_CATALOG, start=1):
            for k in required_keys:
                self.assertIn(k, item, f"Formula #{idx} ({item.get('name')}) missing key: {k}")
                val = item[k]
                self.assertIsInstance(val, str, f"Formula #{idx} ({item.get('name')}) key {k} must be str")
                self.assertTrue(len(val.strip()) > 0, f"Formula #{idx} ({item.get('name')}) key {k} is empty")

            # Check that formulas start with '=' in examples
            self.assertTrue(
                item["example"].startswith("="),
                f"Formula #{idx} ({item['name']}) example must start with '=': {item['example']}"
            )
            categories.add(item["category"])

        # Check all 6 domain categories are present
        expected_categories = {"lookup", "aggregation", "text", "logic", "dynamic", "financial"}
        self.assertTrue(expected_categories.issubset(categories))

    def test_filter_and_search_logic(self):
        """Test category filtering and multi-field keyword searching."""
        # Category filter
        lookup_items = get_formula_catalog(category="lookup")
        self.assertGreaterEqual(len(lookup_items), 5)
        for item in lookup_items:
            self.assertEqual(item["category"], "lookup")

        agg_items = get_formula_catalog(category="aggregation")
        self.assertGreaterEqual(len(agg_items), 5)

        # Keyword search by function name
        xlookup_res = get_formula_catalog(search="XLOOKUP")
        self.assertTrue(any(f["name"] == "XLOOKUP" for f in xlookup_res))

        # Keyword search by scenario/rationale text
        vlookup_comp = get_formula_catalog(search="VLOOKUP")
        self.assertGreater(len(vlookup_comp), 0)

        # Non-matching search
        empty_res = get_formula_catalog(search="NON_EXISTENT_XYZ_12345")
        self.assertEqual(len(empty_res), 0)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed")
    def test_generate_formula_playbook_xlsx(self):
        """Test generation of dual-sheet workbook (.xlsx) with live formulas and styling."""
        out_path = os.path.join(self.tmp_dir.name, "Test_Playbook.xlsx")
        ok, res_path = generate_formula_playbook_xlsx(out_path)

        self.assertTrue(ok)
        self.assertTrue(os.path.isfile(res_path))

        wb = openpyxl.load_workbook(res_path, data_only=False)
        self.assertIn("Formula_Playbook", wb.sheetnames)
        self.assertIn("Live_Interactive_Formulas", wb.sheetnames)

        # Validate Sheet 1 (Reference Guide)
        ws_guide = wb["Formula_Playbook"]
        self.assertGreaterEqual(ws_guide.max_row, 35)
        self.assertGreaterEqual(ws_guide.max_column, 7)
        self.assertEqual(ws_guide.freeze_panes, "A4")

        # Check headers on row 3
        headers = [ws_guide.cell(row=3, column=col).value for col in range(1, 8)]
        self.assertIn("STT", headers[0])
        self.assertIn("Tên Hàm", headers[2])
        self.assertIn("Nhóm Nghiệp Vụ", headers[1])
        self.assertIn("Cú Pháp", headers[3])

        # Validate Sheet 2 (Live Demo)
        ws_live = wb["Live_Interactive_Formulas"]
        self.assertEqual(ws_live.freeze_panes, "A5")

        # Collect all formulas in Sheet 2
        formulas_in_demo = []
        for row in ws_live.iter_rows(values_only=False):
            for cell in row:
                if str(cell.value or "").startswith("="):
                    formulas_in_demo.append(str(cell.value))

        self.assertGreaterEqual(len(formulas_in_demo), 6)
        all_formulas_str = " ".join(formulas_in_demo)
        self.assertIn("SUMIFS", all_formulas_str)
        self.assertIn("COUNTIFS", all_formulas_str)
        self.assertIn("AVERAGEIFS", all_formulas_str)
        self.assertIn("XLOOKUP", all_formulas_str)
        self.assertIn("IFS", all_formulas_str)

        wb.close()

    def test_mcp_server_formula_helper_dispatch(self):
        """Test MCP tool dispatcher for xlsx_formula_helper."""
        res = execute_xlsx_formula_helper({"category": "lookup", "search": "XLOOKUP"})
        self.assertTrue(res.get("success"))
        self.assertEqual(res.get("category"), "lookup")
        self.assertGreaterEqual(res.get("total_count", 0), 1)
        self.assertTrue(any(f["name"] == "XLOOKUP" for f in res["formulas"]))

    def test_cli_subcommand_dispatch(self):
        """Test CLI command handler cmd_xlsx_formulas."""
        # 1. Test listing output
        args_list = argparse.Namespace(
            category="logic",
            search=None,
            export=None,
            json=False,
        )
        saved_stdout = sys.stdout
        sys.stdout = io.StringIO()
        try:
            exit_code = cmd_xlsx_formulas(args_list)
            out_str = sys.stdout.getvalue()
        finally:
            sys.stdout = saved_stdout

        self.assertEqual(exit_code, 0)
        self.assertIn("EXCEL FORMULA PLAYBOOK", out_str)
        self.assertIn("IFS", out_str)

        # 2. Test JSON output
        args_json = argparse.Namespace(
            category="all",
            search="INDEX",
            export=None,
            json=True,
        )
        sys.stdout = io.StringIO()
        try:
            exit_code = cmd_xlsx_formulas(args_json)
            out_json_str = sys.stdout.getvalue()
        finally:
            sys.stdout = saved_stdout

        self.assertEqual(exit_code, 0)
        self.assertIn('"name": "INDEX / MATCH"', out_json_str)

        # 3. Test Export via CLI
        cli_export_file = os.path.join(self.tmp_dir.name, "CLI_Export.xlsx")
        args_export = argparse.Namespace(
            category="all",
            search=None,
            export=cli_export_file,
            json=False,
        )
        sys.stdout = io.StringIO()
        try:
            exit_code = cmd_xlsx_formulas(args_export)
        finally:
            sys.stdout = saved_stdout

        self.assertEqual(exit_code, 0)
        self.assertTrue(os.path.isfile(cli_export_file))


if __name__ == "__main__":
    unittest.main()
