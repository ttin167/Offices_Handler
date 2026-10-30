"""
tests/test_docx_lists.py — Comprehensive Unit & Integration Tests for OpenXML Lists Engine.
Tests cover:
1. AbstractNum registration (2001 decimal, 2002 bullet) and idempotency.
2. Indentation-aware Markdown list line parsing (levels 0, 1, 2).
3. Checkbox task list parsing (unchecked [ ], checked [x]/[X]).
4. Dynamic numId allocation and clean restarts at 1.0 (startOverride=1).
5. Smart continuation (sub-items do not clobber parent numbering).
6. Strict guard against auto-numbering normal text.
7. OpenXML indentation & hanging invariants (<w:ind w:left="..." w:hanging="360"/>).
8. Markdown-to-DOCX list compilation across sections and headings.
9. Table cell multiline list parsing and ERR_DOCX_001 Last Paragraph Rule.
10. docx_writer AST generation with list properties.
"""

import os
import unittest
import docx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

from core.docx.lists import (
    ListContextTracker,
    ParsedListItem,
    apply_list_formatting,
    ensure_numbering_definitions,
    ABSTRACT_NUM_DECIMAL_ID,
    ABSTRACT_NUM_BULLET_ID,
)
from core.converter.markdown import decoupled_to_docx
from core.docx.writer import write_docx


class TestDocxListsEngine(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.tracker = ListContextTracker(self.doc)

    def test_01_ensure_numbering_definitions_creates_abstract_nums(self):
        """AbstractNum 2001 and 2002 must exist in numbering.xml."""
        ensure_numbering_definitions(self.doc)
        n_el = self.doc.part.numbering_part.element
        abstract_ids = [
            int(el.get(qn("w:abstractNumId")))
            for el in n_el.findall(qn("w:abstractNum"))
            if el.get(qn("w:abstractNumId")) is not None
        ]
        self.assertIn(ABSTRACT_NUM_DECIMAL_ID, abstract_ids)
        self.assertIn(ABSTRACT_NUM_BULLET_ID, abstract_ids)

    def test_02_ensure_numbering_definitions_is_idempotent(self):
        """Calling ensure_numbering_definitions repeatedly must not duplicate IDs."""
        ensure_numbering_definitions(self.doc)
        ensure_numbering_definitions(self.doc)
        n_el = self.doc.part.numbering_part.element
        abstract_ids = [
            int(el.get(qn("w:abstractNumId")))
            for el in n_el.findall(qn("w:abstractNum"))
            if el.get(qn("w:abstractNumId")) is not None
        ]
        self.assertEqual(abstract_ids.count(ABSTRACT_NUM_DECIMAL_ID), 1)
        self.assertEqual(abstract_ids.count(ABSTRACT_NUM_BULLET_ID), 1)

    def test_03_parse_ordered_level_0(self):
        """'1. Item' or '2) Item' with 0-1 leading spaces is level 0 ordered."""
        item = self.tracker.parse_markdown_line("1. First item")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "ordered")
        self.assertEqual(item.level, 0)
        self.assertEqual(item.text, "First item")
        self.assertEqual(item.raw_number, 1)

        item2 = self.tracker.parse_markdown_line("2) Second item")
        self.assertIsNotNone(item2)
        self.assertEqual(item2.item_type, "ordered")
        self.assertEqual(item2.level, 0)
        self.assertEqual(item2.text, "Second item")
        self.assertEqual(item2.raw_number, 2)

    def test_04_parse_ordered_level_1(self):
        """'  1. Sub item' with 2-3 spaces is level 1 ordered."""
        item = self.tracker.parse_markdown_line("  1. Sub item")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "ordered")
        self.assertEqual(item.level, 1)
        self.assertEqual(item.text, "Sub item")

    def test_05_parse_ordered_level_2(self):
        """'    1. Deep sub item' with >= 4 spaces is level 2 ordered."""
        item = self.tracker.parse_markdown_line("    1. Deep sub item")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "ordered")
        self.assertEqual(item.level, 2)
        self.assertEqual(item.text, "Deep sub item")

    def test_06_parse_bullet_level_0(self):
        """'- Bullet' or '* Bullet' or '+ Bullet' is level 0 bullet."""
        for prefix in ["- ", "* ", "+ "]:
            item = self.tracker.parse_markdown_line(f"{prefix}Disc bullet")
            self.assertIsNotNone(item)
            self.assertEqual(item.item_type, "bullet")
            self.assertEqual(item.level, 0)
            self.assertEqual(item.text, "Disc bullet")

    def test_07_parse_bullet_level_1(self):
        """'  - Circle bullet' is level 1 bullet."""
        item = self.tracker.parse_markdown_line("  - Circle bullet")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "bullet")
        self.assertEqual(item.level, 1)
        self.assertEqual(item.text, "Circle bullet")

    def test_08_parse_bullet_level_2(self):
        """'    - Square bullet' is level 2 bullet."""
        item = self.tracker.parse_markdown_line("    - Square bullet")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "bullet")
        self.assertEqual(item.level, 2)
        self.assertEqual(item.text, "Square bullet")

    def test_09_parse_checkbox_unchecked(self):
        """'- [ ] Task item' is unchecked checkbox."""
        item = self.tracker.parse_markdown_line("- [ ] Pending task")
        self.assertIsNotNone(item)
        self.assertEqual(item.item_type, "checkbox")
        self.assertEqual(item.level, 0)
        self.assertFalse(item.is_checked)
        self.assertEqual(item.text, "Pending task")

    def test_10_parse_checkbox_checked(self):
        """'- [x] Task done' or '* [X] Task done' is checked checkbox."""
        for line in ["- [x] Completed task", "* [X] Another completed"]:
            item = self.tracker.parse_markdown_line(line)
            self.assertIsNotNone(item)
            self.assertEqual(item.item_type, "checkbox")
            self.assertTrue(item.is_checked)

    def test_11_parse_non_list_lines_returns_none(self):
        """Normal paragraphs, headers, and code delimiters return None."""
        self.assertIsNone(self.tracker.parse_markdown_line("This is plain paragraph."))
        self.assertIsNone(self.tracker.parse_markdown_line("# Heading 1"))
        self.assertIsNone(self.tracker.parse_markdown_line("```python"))
        self.assertIsNone(self.tracker.parse_markdown_line(""))
        self.assertIsNone(self.tracker.parse_markdown_line("    "))

    def test_12_ordered_num_id_allocation_and_clean_restart(self):
        """New ordered lists allocate new numId elements with startOverride=1."""
        item1 = self.tracker.parse_markdown_line("1. First list item")
        num_id_1 = self.tracker.get_num_id_for_item(item1)
        self.assertIsNotNone(num_id_1)

        # Second continuous item retains same num_id
        item2 = self.tracker.parse_markdown_line("2. Second list item")
        num_id_2 = self.tracker.get_num_id_for_item(item2)
        self.assertEqual(num_id_1, num_id_2)

        # Explicit reset (e.g. intervening heading or paragraph)
        self.tracker.reset_ordered_list()
        item3 = self.tracker.parse_markdown_line("1. New list item")
        num_id_3 = self.tracker.get_num_id_for_item(item3)
        self.assertIsNotNone(num_id_3)
        self.assertNotEqual(num_id_1, num_id_3)

    def test_13_smart_continuation_with_interleaved_bullets(self):
        """Sub-bullets interleaved between ordered steps must NOT reset parent numId."""
        item1 = self.tracker.parse_markdown_line("1. Step 1")
        num_id_1 = self.tracker.get_num_id_for_item(item1)

        sub_bullet = self.tracker.parse_markdown_line("  - Sub note for step 1")
        num_id_sub = self.tracker.get_num_id_for_item(sub_bullet)
        self.assertNotEqual(num_id_1, num_id_sub)

        # Parent continues
        item2 = self.tracker.parse_markdown_line("2. Step 2")
        num_id_2 = self.tracker.get_num_id_for_item(item2)
        self.assertEqual(num_id_1, num_id_2)

    def test_14_apply_list_formatting_xml_structure(self):
        """apply_list_formatting must insert <w:numPr> and <w:ind> elements."""
        p = self.doc.add_paragraph()
        p.add_run("Sample item")
        apply_list_formatting(p, item_type="ordered", level=0, num_id=3001)

        pPr = p._p.pPr
        numPr = pPr.find(qn("w:numPr"))
        self.assertIsNotNone(numPr)
        ilvl = numPr.find(qn("w:ilvl"))
        self.assertEqual(ilvl.get(qn("w:val")), "0")
        numId = numPr.find(qn("w:numId"))
        self.assertEqual(numId.get(qn("w:val")), "3001")

        ind = pPr.find(qn("w:ind"))
        self.assertIsNotNone(ind)
        self.assertEqual(ind.get(qn("w:left")), "720")
        self.assertEqual(ind.get(qn("w:hanging")), "360")

    def test_15_apply_list_formatting_nesting_levels(self):
        """Level 1 must have left=1440, Level 2 must have left=2160."""
        p1 = self.doc.add_paragraph()
        apply_list_formatting(p1, item_type="bullet", level=1, num_id=3002)
        ind1 = p1._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind1.get(qn("w:left")), "1440")

        p2 = self.doc.add_paragraph()
        apply_list_formatting(p2, item_type="bullet", level=2, num_id=3002)
        ind2 = p2._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind2.get(qn("w:left")), "2160")

    def test_16_apply_list_formatting_checkbox_glyphs(self):
        """Checkboxes prepend ballot box glyphs and set hanging indentation."""
        p_uncheck = self.doc.add_paragraph()
        p_uncheck.add_run("Unchecked task")
        apply_list_formatting(p_uncheck, item_type="checkbox", level=0, is_checked=False)
        self.assertTrue(p_uncheck.runs[0].text.startswith("☐"))

        p_check = self.doc.add_paragraph()
        p_check.add_run("Checked task")
        apply_list_formatting(p_check, item_type="checkbox", level=0, is_checked=True)
        self.assertTrue(p_check.runs[0].text.startswith("☑"))

    def test_17_markdown_to_docx_ordered_lists_have_distinct_num_ids(self):
        """Separate markdown sections with ordered lists must receive distinct numIds."""
        md_text = (
            "# Section 1\n"
            "1. Apple\n"
            "2. Banana\n"
            "\n"
            "# Section 2\n"
            "1. Orange\n"
            "2. Grape\n"
        )
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        os.makedirs(out_dir, exist_ok=True)
        md_file = os.path.join(out_dir, "test_lists_sections.md")
        docx_file = os.path.join(out_dir, "test_lists_sections.docx")

        with open(md_file, "w", encoding="utf-8") as f:
            f.write(md_text)

        decoupled_to_docx(md_file, docx_file)

        doc_read = Document(docx_file)
        list_paragraphs = [
            p for p in doc_read.paragraphs
            if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        ]
        self.assertEqual(len(list_paragraphs), 4)

        # Section 1 numId
        sec1_num_id_1 = list_paragraphs[0]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        sec1_num_id_2 = list_paragraphs[1]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        self.assertEqual(sec1_num_id_1, sec1_num_id_2)

        # Section 2 numId
        sec2_num_id_1 = list_paragraphs[2]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        sec2_num_id_2 = list_paragraphs[3]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        self.assertEqual(sec2_num_id_1, sec2_num_id_2)

        # Must restart: Section 1 numId != Section 2 numId
        self.assertNotEqual(sec1_num_id_1, sec2_num_id_1)

    def test_18_markdown_to_docx_plain_text_has_no_numbering(self):
        """Zero auto-numbering invariant: normal paragraphs must NOT contain numPr."""
        md_text = (
            "# Main Title\n"
            "This is normal paragraph one.\n"
            "This is normal paragraph two.\n"
            "1. This is a list item.\n"
            "This is normal paragraph three.\n"
        )
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        md_file = os.path.join(out_dir, "test_lists_plain.md")
        docx_file = os.path.join(out_dir, "test_lists_plain.docx")

        with open(md_file, "w", encoding="utf-8") as f:
            f.write(md_text)

        decoupled_to_docx(md_file, docx_file)

        doc_read = Document(docx_file)
        numbered_count = 0
        for p in doc_read.paragraphs:
            if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None:
                numbered_count += 1
        # Exactly one numbered paragraph
        self.assertEqual(numbered_count, 1)

    def test_19_markdown_to_docx_table_cell_multiline_lists(self):
        """Table cells with multiline bullet items must format as proper list paragraphs."""
        md_text = (
            "| Feature | Checklist |\n"
            "|---|---|\n"
            "| Core | - Auth<br>- Profile<br>- Settings |\n"
            "| API | 1. Login<br>2. Logout |\n"
        )
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        md_file = os.path.join(out_dir, "test_lists_table.md")
        docx_file = os.path.join(out_dir, "test_lists_table.docx")

        with open(md_file, "w", encoding="utf-8") as f:
            f.write(md_text)

        decoupled_to_docx(md_file, docx_file)

        doc_read = Document(docx_file)
        self.assertTrue(len(doc_read.tables) > 0)
        table = doc_read.tables[0]

        # Row 1, Cell 1 has 3 bullet items
        cell_bullets = table.cell(1, 1)
        self.assertEqual(len(cell_bullets.paragraphs), 3)
        for cp in cell_bullets.paragraphs:
            self.assertIsNotNone(cp._p.pPr.find(qn("w:numPr")))

        # Row 2, Cell 1 has 2 numbered items
        cell_numbers = table.cell(2, 1)
        self.assertEqual(len(cell_numbers.paragraphs), 2)
        for cp in cell_numbers.paragraphs:
            self.assertIsNotNone(cp._p.pPr.find(qn("w:numPr")))

    def test_20_docx_writer_ast_generates_lists(self):
        """docx_writer.write_docx must format paragraphs with list_type correctly."""
        snapshot = {
            "body": [
                {
                    "type": "paragraph",
                    "text": "Heading Item",
                    "style_name": "Heading 1"
                },
                {
                    "type": "paragraph",
                    "text": "First numbered step",
                    "list_type": "ordered",
                    "list_level": 0
                },
                {
                    "type": "paragraph",
                    "text": "Sub-step letter a",
                    "list_type": "ordered",
                    "list_level": 1
                },
                {
                    "type": "paragraph",
                    "text": "Action bullet item",
                    "list_type": "bullet",
                    "list_level": 0
                },
                {
                    "type": "paragraph",
                    "text": "Task to complete",
                    "list_type": "checkbox",
                    "is_checked": True
                }
            ]
        }
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        out_docx = os.path.join(out_dir, "test_writer_lists.docx")
        write_docx(snapshot, out_docx)

        doc_read = Document(out_docx)
        self.assertEqual(len(doc_read.paragraphs), 5)

        # Check list item 1 (ordered lvl 0)
        p1 = doc_read.paragraphs[1]
        self.assertIsNotNone(p1._p.pPr.find(qn("w:numPr")))
        self.assertEqual(p1._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl")).get(qn("w:val")), "0")

        # Check list item 2 (ordered lvl 1)
        p2 = doc_read.paragraphs[2]
        self.assertIsNotNone(p2._p.pPr.find(qn("w:numPr")))
        self.assertEqual(p2._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl")).get(qn("w:val")), "1")

        # Check list item 3 (bullet lvl 0)
        p3 = doc_read.paragraphs[3]
        self.assertIsNotNone(p3._p.pPr.find(qn("w:numPr")))

        # Check list item 4 (checkbox checked)
        p4 = doc_read.paragraphs[4]
        self.assertTrue(p4.runs[0].text.startswith("☑"))

    def test_21_docx_writer_list_restart_flag(self):
        """list_restart=True in para_data must trigger a fresh numId."""
        snapshot = {
            "body": [
                {"type": "paragraph", "text": "Step 1", "list_type": "ordered"},
                {"type": "paragraph", "text": "Step 2", "list_type": "ordered"},
                {"type": "paragraph", "text": "Step 1 restarted", "list_type": "ordered", "list_restart": True},
            ]
        }
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        out_docx = os.path.join(out_dir, "test_writer_restart.docx")
        write_docx(snapshot, out_docx)

        doc_read = Document(out_docx)
        id1 = doc_read.paragraphs[0]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        id2 = doc_read.paragraphs[1]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        id3 = doc_read.paragraphs[2]._p.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))

        self.assertEqual(id1, id2)
        self.assertNotEqual(id1, id3)

    def test_22_table_cell_empty_has_last_paragraph(self):
        """Empty table cells must retain at least one <w:p> (ERR_DOCX_001)."""
        snapshot = {
            "body": [
                {
                    "type": "table",
                    "rows": [
                        {"cells": [{"text": "Header"}]},
                        {"cells": [{"text": ""}]}
                    ]
                }
            ]
        }
        out_dir = os.path.join(os.getcwd(), ".web_outputs")
        out_docx = os.path.join(out_dir, "test_table_last_p.docx")
        write_docx(snapshot, out_docx)

        doc_read = Document(out_docx)
        tc_empty = doc_read.tables[0].rows[1].cells[0]._tc
        self.assertTrue(len(tc_empty.findall(qn("w:p"))) >= 1)


if __name__ == "__main__":
    unittest.main()
