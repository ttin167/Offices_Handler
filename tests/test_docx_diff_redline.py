"""
tests/test_docx_diff_redline.py — Unit test suite for Legal Redline Diff Engine (ERR_DOCX_008).
"""

from __future__ import annotations

import os
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import pytest
from docx import Document
from docx.oxml.ns import qn

from core.docx.diff_engine import (
    RedlineDiffEngine,
    TrackChangeIdGenerator,
    tokenize_text,
    diff_token_sequences,
    create_ins_element,
    create_del_element,
    compare_docx_documents,
)
from tools.unified_qa_diagnostic import DocxDiagnosticEngine


def test_tokenize_text_roundtrip():
    text = "Hợp đồng lao động số 123/2026/HĐLĐ ký ngày 19/10/2026 giữa Công ty A và Ông B."
    tokens = tokenize_text(text)
    assert "".join(tokens) == text
    assert "Hợp" in tokens
    assert " " in tokens
    assert "/" in tokens


def test_diff_token_sequences():
    tokens_a = tokenize_text("Bên A trả 10 triệu đồng mỗi tháng.")
    tokens_b = tokenize_text("Bên A trả 15 triệu đồng mỗi quý.")
    diffs = diff_token_sequences(tokens_a, tokens_b)

    tags = [d[0] for d in diffs]
    assert "equal" in tags
    assert "delete" in tags
    assert "insert" in tags

    deleted_text = "".join(chunk for tag, chunk in diffs if tag == "delete")
    inserted_text = "".join(chunk for tag, chunk in diffs if tag == "insert")
    assert "10" in deleted_text
    assert "15" in inserted_text
    assert "tháng" in deleted_text
    assert "quý" in inserted_text


def test_create_ins_element_schema():
    ins = create_ins_element(
        "nội dung mới", change_id=1, author="Lawyer AI", date_str="2026-10-19T08:00:00Z"
    )
    assert ins.tag.endswith("ins")
    assert ins.get(qn("w:id")) == "1"
    assert ins.get(qn("w:author")) == "Lawyer AI"
    assert ins.get(qn("w:date")) == "2026-10-19T08:00:00Z"
    # Contains <w:t>
    t_elem = ins.find(f".//{qn('w:t')}")
    assert t_elem is not None
    assert t_elem.text == "nội dung mới"


def test_create_del_element_schema_err_docx_008():
    del_elem = create_del_element(
        "nội dung cũ", change_id=2, author="Lawyer AI", date_str="2026-10-19T08:00:00Z"
    )
    assert del_elem.tag.endswith("del")
    assert del_elem.get(qn("w:id")) == "2"
    assert del_elem.get(qn("w:author")) == "Lawyer AI"
    assert del_elem.get(qn("w:date")) == "2026-10-19T08:00:00Z"

    # Crucial ERR_DOCX_008 invariant: <w:del> must contain <w:delText>, NOT <w:t>
    del_text_elem = del_elem.find(f".//{qn('w:delText')}")
    assert del_text_elem is not None
    assert del_text_elem.text == "nội dung cũ"
    assert del_elem.find(f".//{qn('w:t')}") is None


def test_track_change_id_generator():
    gen = TrackChangeIdGenerator(start_id=10)
    assert gen.next_id() == 10
    assert gen.next_id() == 11
    assert gen.next_id() == 12


def test_compare_docx_identical_documents(tmp_path):
    doc_a = Document()
    doc_a.add_heading("Điều khoản chung", level=1)
    doc_a.add_paragraph("Hai bên cam kết thực hiện đúng mọi điều khoản đã ký kết.")
    path_a = str(tmp_path / "doc_a.docx")
    doc_a.save(path_a)

    path_diff = str(tmp_path / "diff_identical.docx")
    stats = compare_docx_documents(path_a, path_a, path_diff, author="AI Auditor")

    assert stats["insertions_count"] == 0
    assert stats["deletions_count"] == 0
    assert os.path.exists(path_diff)

    # Document should open cleanly
    res_doc = Document(path_diff)
    assert len(res_doc.paragraphs) >= 2


def test_compare_docx_with_modifications(tmp_path):
    doc_a = Document()
    doc_a.add_paragraph("Thời hạn hợp đồng là 12 tháng kể từ ngày ký.")
    doc_a.add_paragraph("Khoản tiền phạt vi phạm là 5% giá trị hợp đồng.")
    path_a = str(tmp_path / "contract_v1.docx")
    doc_a.save(path_a)

    doc_b = Document()
    doc_b.add_paragraph("Thời hạn hợp đồng là 24 tháng kể từ ngày ký kết.")
    doc_b.add_paragraph("Khoản tiền phạt vi phạm là 8% giá trị hợp đồng.")
    path_b = str(tmp_path / "contract_v2.docx")
    doc_b.save(path_b)

    path_diff = str(tmp_path / "contract_diff.docx")
    stats = compare_docx_documents(path_a, path_b, path_diff, author="AI Legal Reviewer")

    assert stats["insertions_count"] > 0
    assert stats["deletions_count"] > 0
    assert stats["paragraphs_compared"] == 2
    assert os.path.exists(path_diff)

    # Open XML validation
    diff_doc = Document(path_diff)
    body_xml = diff_doc._body._element.xml
    assert "<w:ins" in body_xml
    assert "<w:del" in body_xml
    assert "<w:delText" in body_xml


def test_compare_docx_with_tables(tmp_path):
    doc_a = Document()
    t_a = doc_a.add_table(rows=2, cols=2)
    t_a.rows[0].cells[0].paragraphs[0].text = "STT"
    t_a.rows[0].cells[1].paragraphs[0].text = "Tên Dịch Vụ"
    t_a.rows[1].cells[0].paragraphs[0].text = "01"
    t_a.rows[1].cells[1].paragraphs[0].text = "Dịch vụ Bảo trì Tiêu chuẩn"
    path_a = str(tmp_path / "table_v1.docx")
    doc_a.save(path_a)

    doc_b = Document()
    t_b = doc_b.add_table(rows=2, cols=2)
    t_b.rows[0].cells[0].paragraphs[0].text = "STT"
    t_b.rows[0].cells[1].paragraphs[0].text = "Tên Dịch Vụ"
    t_b.rows[1].cells[0].paragraphs[0].text = "01"
    t_b.rows[1].cells[1].paragraphs[0].text = "Dịch vụ Bảo trì Nâng cao 24/7"
    path_b = str(tmp_path / "table_v2.docx")
    doc_b.save(path_b)

    path_diff = str(tmp_path / "table_diff.docx")
    stats = compare_docx_documents(path_a, path_b, path_diff)

    assert stats["tables_compared"] == 1
    assert stats["insertions_count"] > 0
    assert stats["deletions_count"] > 0

    diff_doc = Document(path_diff)
    assert len(diff_doc.tables) == 1
    cell = diff_doc.tables[0].rows[1].cells[1]
    # Invariant: Cell ends with <w:p>
    assert len(cell._tc.findall(qn("w:p"))) >= 1


def test_compare_docx_with_comments_injection(tmp_path):
    doc_a = Document()
    doc_a.add_paragraph("Mức bồi thường thiệt hại tối đa là 100.000.000 VNĐ.")
    path_a = str(tmp_path / "liability_v1.docx")
    doc_a.save(path_a)

    doc_b = Document()
    doc_b.add_paragraph("Mức bồi thường thiệt hại tối đa không vượt quá 50.000.000 VNĐ.")
    path_b = str(tmp_path / "liability_v2.docx")
    doc_b.save(path_b)

    path_diff = str(tmp_path / "liability_diff.docx")
    comments = [{"text": "Lưu ý pháp lý: Cần đối chiếu với Điều 360 Bộ luật Dân sự 2015."}]
    stats = compare_docx_documents(path_a, path_b, path_diff, comments=comments)

    assert stats["comments_count"] == 1

    # Verify word/comments.xml in the ZIP package
    with zipfile.ZipFile(path_diff, "r") as zf:
        namelist = zf.namelist()
        assert "word/comments.xml" in namelist
        comments_xml = zf.read("word/comments.xml").decode("utf-8")
        assert "Bộ luật Dân sự" in comments_xml
        assert "w:comment" in comments_xml


def test_unified_qa_diagnostic_validates_clean_err_docx_008(tmp_path):
    doc_a = Document()
    doc_a.add_paragraph("Chính sách thanh toán: chuyển khoản 30 ngày.")
    path_a = str(tmp_path / "policy_v1.docx")
    doc_a.save(path_a)

    doc_b = Document()
    doc_b.add_paragraph("Chính sách thanh toán: chuyển khoản ngay trong 7 ngày làm việc.")
    path_b = str(tmp_path / "policy_v2.docx")
    doc_b.save(path_b)

    path_diff = str(tmp_path / "policy_diff.docx")
    compare_docx_documents(doc_a_path=path_a, doc_b_path=path_b, output_docx_path=path_diff)

    # Run DocxDiagnosticEngine to ensure 0 ERR_DOCX_008 errors
    issues, metrics = DocxDiagnosticEngine.diagnose(path_diff)
    docx_008_issues = [i for i in issues if i.code == "ERR_DOCX_008"]
    assert len(docx_008_issues) == 0, f"Found unexpected ERR_DOCX_008 issues: {[i.message for i in docx_008_issues]}"
