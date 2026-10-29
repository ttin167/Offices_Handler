# -*- coding: utf-8 -*-
"""
core/xlsx/reporting.py — Two-Way Spreadsheet Reporting & Multi-Workbook Consolidation Engine (v4.6).

Capabilities:
  1. excel_to_brief: Converts multi-sheet Excel workbooks (.xlsx) into publication-grade
     Word (.docx), Markdown (.md), or PDF (.pdf) reports with automatic KPI aggregation
     and strict OpenXML table invariants (<w:cantSplit/>, <w:tblHeader/>, <w:vAlign w:val="center"/>).
  2. Dual-Workbook Formula Evaluator Fallback (ERR_XLSX_015 / Edge Case #3):
     When reading newly generated openpyxl workbooks where data_only=True returns None
     for formula cells, automatically evaluates SUM, AVERAGE, COUNT, COUNTA, MIN, MAX,
     and cell arithmetic expressions in-memory so summary rows are never blank.
  3. consolidate_workbooks: Merges multiple departmental/branch Excel files sharing a template
     into a Master Workbook with header schema validation, leading-zero preservation (ERR_XLSX_007),
     and live summary formulas.
"""

from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from .writer import apply_worksheet_print_setup, sanitize_xml_string

CELL_REF_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
RANGE_REF_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+):\$?([A-Za-z]{1,3})\$?(\d+)$")
FUNC_CALL_RE = re.compile(
    r"^=\s*(?:_xlfn\.)?(SUM|AVERAGE|COUNT|COUNTA|MIN|MAX)\s*\((.+)\)\s*$",
    re.IGNORECASE,
)
ARITH_EXPR_RE = re.compile(r"^=\s*([A-Za-z0-9_$\.\+\-\*\/\(\)\s]+)$")


def _to_number(val: Any) -> Optional[float]:
    """Convert value to float if numeric, else None."""
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        cleaned = val.strip().replace(",", "")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _resolve_cell_value(
    ws_val: Any,
    ws_formula: Any,
    row: int,
    col: int,
    _depth: int = 0,
) -> Any:
    """
    Resolves cell value at (row, col) with Dual-Workbook Formula Fallback (ERR_XLSX_015).
    """
    if _depth > 6:
        return None

    v = ws_val.cell(row=row, column=col).value
    if v is not None:
        return v

    f_val = ws_formula.cell(row=row, column=col).value
    if isinstance(f_val, str) and f_val.strip().startswith("="):
        return _evaluate_basic_excel_formula(f_val.strip(), ws_val, ws_formula, _depth=_depth + 1)
    return f_val


def _expand_range_values(
    arg_str: str,
    ws_val: Any,
    ws_formula: Any,
    _depth: int = 0,
) -> List[Any]:
    """Expands a formula argument list (e.g. 'B2:B5, D2') into resolved cell values."""
    values: List[Any] = []
    for token in arg_str.split(","):
        part = token.strip()
        if not part:
            continue
        if "!" in part:
            part = part.split("!", 1)[1].strip()

        m_range = RANGE_REF_RE.match(part)
        if m_range:
            c1 = column_index_from_string(m_range.group(1).upper())
            r1 = int(m_range.group(2))
            c2 = column_index_from_string(m_range.group(3).upper())
            r2 = int(m_range.group(4))
            for r in range(min(r1, r2), max(r1, r2) + 1):
                for c in range(min(c1, c2), max(c1, c2) + 1):
                    values.append(_resolve_cell_value(ws_val, ws_formula, r, c, _depth=_depth))
            continue

        m_cell = CELL_REF_RE.match(part)
        if m_cell:
            c = column_index_from_string(m_cell.group(1).upper())
            r = int(m_cell.group(2))
            values.append(_resolve_cell_value(ws_val, ws_formula, r, c, _depth=_depth))
            continue

        num = _to_number(part)
        if num is not None:
            values.append(num)
        else:
            values.append(part.strip('"\''))
    return values


def _evaluate_basic_excel_formula(
    formula_str: str,
    ws_val: Any,
    ws_formula: Any,
    _depth: int = 0,
) -> Any:
    """
    Evaluates basic Excel formulas (SUM, AVERAGE, COUNT, COUNTA, MIN, MAX, and basic
    cell arithmetic) when openpyxl data_only=True returns None on un-cached workbooks (ERR_XLSX_015).
    """
    clean_f = formula_str.strip()
    m_func = FUNC_CALL_RE.match(clean_f)
    if m_func:
        func_name = m_func.group(1).upper()
        inner_args = m_func.group(2)
        raw_vals = _expand_range_values(inner_args, ws_val, ws_formula, _depth=_depth)
        nums = [n for n in (_to_number(v) for v in raw_vals) if n is not None]

        if func_name == "SUM":
            res = sum(nums) if nums else 0.0
            return int(res) if res == int(res) else round(res, 2)
        if func_name == "AVERAGE":
            if not nums:
                return 0.0
            res = sum(nums) / len(nums)
            return int(res) if res == int(res) else round(res, 2)
        if func_name == "COUNT":
            return len(nums)
        if func_name == "COUNTA":
            return sum(1 for v in raw_vals if v is not None and str(v).strip() != "")
        if func_name == "MIN":
            if not nums:
                return 0
            res = min(nums)
            return int(res) if res == int(res) else round(res, 2)
        if func_name == "MAX":
            if not nums:
                return 0
            res = max(nums)
            return int(res) if res == int(res) else round(res, 2)

    # Handle simple cell arithmetic expressions like =B2+C2 or =B2*C2
    if ARITH_EXPR_RE.match(clean_f) and "(" not in clean_f[1:]:
        expr = clean_f[1:].strip()

        def _replace_ref(match: re.Match) -> str:
            token = match.group(0)
            m_c = CELL_REF_RE.match(token)
            if not m_c:
                return token
            c_idx = column_index_from_string(m_c.group(1).upper())
            r_idx = int(m_c.group(2))
            resolved = _resolve_cell_value(ws_val, ws_formula, r_idx, c_idx, _depth=_depth)
            num_val = _to_number(resolved)
            return str(num_val if num_val is not None else 0.0)

        substituted = re.sub(r"\$?[A-Za-z]{1,3}\$?\d+", _replace_ref, expr)
        if re.match(r"^[\d\.\+\-\*\/\(\)\s]+$", substituted):
            try:
                val = float(eval(substituted, {"__builtins__": {}}, {}))
                return int(val) if val == int(val) else round(val, 2)
            except Exception:
                pass

    return clean_f


def _format_cell_display(val: Any) -> str:
    """Format a resolved cell value into clean presentation text."""
    if val is None:
        return ""
    if isinstance(val, float):
        if val == int(val):
            return f"{int(val):,}"
        return f"{val:,.2f}"
    if isinstance(val, int) and not isinstance(val, bool):
        return f"{val:,}"
    return sanitize_xml_string(str(val).strip())


def _extract_sheet_matrix_and_kpis(
    ws_val: Any,
    ws_formula: Any,
    max_rows: int = 500,
) -> Dict[str, Any]:
    """Extracts non-empty table rows and computes numeric column KPIs for a worksheet."""
    max_r = min(ws_formula.max_row or 0, max_rows)
    max_c = min(ws_formula.max_column or 0, 50)
    rows: List[List[Any]] = []

    for r in range(1, max_r + 1):
        row_vals = [
            _resolve_cell_value(ws_val, ws_formula, r, c)
            for c in range(1, max_c + 1)
        ]
        if any(v is not None and str(v).strip() != "" for v in row_vals):
            rows.append(row_vals)

    if not rows:
        return {"headers": [], "data_rows": [], "display_rows": [], "kpis": []}

    # Trim trailing empty columns
    last_non_empty_col = 0
    for row in rows:
        for idx, v in enumerate(row):
            if v is not None and str(v).strip() != "":
                last_non_empty_col = max(last_non_empty_col, idx + 1)
    if last_non_empty_col > 0:
        rows = [r[:last_non_empty_col] for r in rows]

    headers = [str(h).strip() if h is not None else f"Col_{i+1}" for i, h in enumerate(rows[0])]
    data_rows = rows[1:]
    display_rows = [[_format_cell_display(c) for c in r] for r in data_rows]

    kpis: List[Dict[str, Any]] = []
    for col_idx, col_name in enumerate(headers):
        col_nums: List[float] = []
        for idx, r in enumerate(data_rows):
            # Exclude bottom total/summary row from KPI double-counting if first cell says Tổng/Total
            first_cell_str = str(r[0] or "").strip().lower() if r else ""
            if idx == len(data_rows) - 1 and any(
                kw in first_cell_str for kw in ("tổng", "total", "sum", "trung bình", "average")
            ):
                continue
            if col_idx < len(r):
                n = _to_number(r[col_idx])
                if n is not None:
                    col_nums.append(n)
        if col_nums:
            total_v = sum(col_nums)
            avg_v = total_v / len(col_nums)
            kpis.append({
                "column": col_name,
                "count": len(col_nums),
                "sum": int(total_v) if total_v == int(total_v) else round(total_v, 2),
                "average": int(avg_v) if avg_v == int(avg_v) else round(avg_v, 2),
                "min": min(col_nums),
                "max": max(col_nums),
            })

    return {
        "headers": headers,
        "data_rows": data_rows,
        "display_rows": display_rows,
        "kpis": kpis,
    }


def _apply_docx_table_invariants(table: Any) -> None:
    """
    Enforces strict OpenXML table invariants:
    - <w:tblHeader/> on the first row
    - <w:cantSplit/> on every row
    - <w:vAlign w:val="center"/> on every cell
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        if row_idx == 0 and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))

        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            if tcPr.find(qn("w:vAlign")) is None:
                vAlign = OxmlElement("w:vAlign")
                vAlign.set(qn("w:val"), "center")
                tcPr.append(vAlign)
            if row_idx == 0 and tcPr.find(qn("w:shd")) is None:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "0F294A")
                tcPr.append(shd)


def excel_to_brief(
    xlsx_path: str,
    output_path: Optional[str] = None,
    target_format: Optional[str] = None,
    title: Optional[str] = None,
    include_kpi_summary: bool = True,
    max_rows_per_sheet: int = 500,
) -> Dict[str, Any]:
    """
    Converts an Excel workbook (.xlsx) into an executive brief report in Word (.docx),
    Markdown (.md), or PDF (.pdf) format while preserving formula totals via Dual-Workbook
    evaluation (ERR_XLSX_015).
    """
    src = os.path.abspath(xlsx_path)
    if not os.path.isfile(src):
        raise FileNotFoundError(f"Excel file not found: {src}")

    inferred_ext = os.path.splitext(output_path)[1].lower().lstrip(".") if output_path else ""
    fmt = (target_format or inferred_ext or "docx").strip().lower().lstrip(".")
    if fmt not in ("docx", "md", "markdown", "pdf"):
        raise ValueError(f"Unsupported target format for excel_to_brief: {target_format}")
    if fmt == "markdown":
        fmt = "md"

    if not output_path:
        output_path = f"{os.path.splitext(src)[0]}_brief.{fmt}"
    dest = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(dest), exist_ok=True)

    # Guard #3 (ERR_XLSX_015): Dual-Workbook load (data_only=True + data_only=False)
    wb_val = openpyxl.load_workbook(src, data_only=True)
    wb_formula = openpyxl.load_workbook(src, data_only=False)

    report_title = title or f"Báo Cáo Tổng Hợp Dữ Liệu — {os.path.splitext(os.path.basename(src))[0]}"
    sheets_summary: List[Dict[str, Any]] = []

    try:
        for sheet_name in wb_formula.sheetnames:
            ws_f = wb_formula[sheet_name]
            ws_v = wb_val[sheet_name]
            if getattr(ws_f, "sheet_state", "visible") != "visible":
                continue
            extracted = _extract_sheet_matrix_and_kpis(ws_v, ws_f, max_rows=max_rows_per_sheet)
            if extracted["headers"]:
                sheets_summary.append({
                    "sheet_name": sheet_name,
                    **extracted,
                })
    finally:
        wb_val.close()
        wb_formula.close()

    # 1. Markdown Generation
    md_lines: List[str] = [f"# {report_title}", ""]
    md_lines.append(f"- **Tệp nguồn**: `{os.path.basename(src)}`")
    md_lines.append(f"- **Số trang tính (Sheets)**: {len(sheets_summary)}")
    md_lines.append("")

    for idx, sdata in enumerate(sheets_summary, start=1):
        md_lines.append(f"## {idx}. Trang tính: {sdata['sheet_name']}")
        md_lines.append("")

        if include_kpi_summary and sdata["kpis"]:
            md_lines.append("### Chỉ Số Thống Kê (KPI Summary)")
            md_lines.append("")
            md_lines.append("| Cột Dữ Liệu | Số Lượng | Tổng Cộng (SUM) | Trung Bình (AVG) | Nhỏ Nhất (MIN) | Lớn Nhất (MAX) |")
            md_lines.append("|---|---|---|---|---|---|")
            for k in sdata["kpis"]:
                md_lines.append(
                    f"| {k['column']} | {k['count']} | {_format_cell_display(k['sum'])} | "
                    f"{_format_cell_display(k['average'])} | {_format_cell_display(k['min'])} | {_format_cell_display(k['max'])} |"
                )
            md_lines.append("")

        md_lines.append("### Bảng Dữ Liệu Chi Tiết")
        md_lines.append("")
        headers = sdata["headers"]
        md_lines.append("| " + " | ".join(h.replace("|", "\\|") for h in headers) + " |")
        md_lines.append("| " + " | ".join("---" for _ in headers) + " |")
        for drow in sdata["display_rows"]:
            padded = (drow + [""] * len(headers))[: len(headers)]
            md_lines.append("| " + " | ".join(c.replace("|", "\\|") for c in padded) + " |")
        md_lines.append("")

    md_content = "\n".join(md_lines)

    if fmt == "md":
        with open(dest, "w", encoding="utf-8") as f:
            f.write(md_content)
        return {
            "success": True,
            "output_path": dest,
            "format": "md",
            "sheets_processed": len(sheets_summary),
            "sheets": [
                {"sheet_name": s["sheet_name"], "rows": len(s["display_rows"]), "kpis": s["kpis"]}
                for s in sheets_summary
            ],
        }

    # 2. Word (.docx) Generation with OpenXML Table Invariants
    import docx
    from docx.shared import Pt, RGBColor

    docx_target = dest if fmt == "docx" else f"{os.path.splitext(dest)[0]}_temp_brief.docx"
    doc = docx.Document()
    doc.add_heading(report_title, level=0)

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        f"Tệp nguồn: {os.path.basename(src)} | Số lượng Sheet: {len(sheets_summary)}"
    )
    meta_run.italic = True

    for idx, sdata in enumerate(sheets_summary, start=1):
        doc.add_heading(f"{idx}. Trang tính: {sdata['sheet_name']}", level=1)

        if include_kpi_summary and sdata["kpis"]:
            doc.add_heading("Chỉ Số Thống Kê (KPI Summary)", level=2)
            kpi_headers = ["Cột Dữ Liệu", "Số Lượng", "Tổng Cộng (SUM)", "Trung Bình (AVG)", "MIN", "MAX"]
            kpi_tbl = doc.add_table(rows=1 + len(sdata["kpis"]), cols=len(kpi_headers))
            kpi_tbl.style = "Table Grid"
            for c_i, h_text in enumerate(kpi_headers):
                p = kpi_tbl.rows[0].cells[c_i].paragraphs[0]
                r = p.add_run(h_text)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.5)
            for r_i, k in enumerate(sdata["kpis"], start=1):
                vals = [
                    str(k["column"]),
                    str(k["count"]),
                    _format_cell_display(k["sum"]),
                    _format_cell_display(k["average"]),
                    _format_cell_display(k["min"]),
                    _format_cell_display(k["max"]),
                ]
                for c_i, val_str in enumerate(vals):
                    p = kpi_tbl.rows[r_i].cells[c_i].paragraphs[0]
                    r = p.add_run(val_str)
                    r.font.size = Pt(9.5)
            _apply_docx_table_invariants(kpi_tbl)
            doc.add_paragraph()

        headers = sdata["headers"]
        drows = sdata["display_rows"]
        if headers:
            doc.add_heading("Bảng Dữ Liệu Chi Tiết", level=2)
            tbl = doc.add_table(rows=1 + len(drows), cols=len(headers))
            tbl.style = "Table Grid"
            for c_i, h_text in enumerate(headers):
                p = tbl.rows[0].cells[c_i].paragraphs[0]
                r = p.add_run(h_text)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9.5)
            for r_i, drow in enumerate(drows, start=1):
                for c_i in range(len(headers)):
                    val_str = drow[c_i] if c_i < len(drow) else ""
                    p = tbl.rows[r_i].cells[c_i].paragraphs[0]
                    r = p.add_run(val_str)
                    r.font.size = Pt(9.5)
            _apply_docx_table_invariants(tbl)
            doc.add_paragraph()

    doc.save(docx_target)

    if fmt == "docx":
        return {
            "success": True,
            "output_path": dest,
            "format": "docx",
            "sheets_processed": len(sheets_summary),
            "sheets": [
                {"sheet_name": s["sheet_name"], "rows": len(s["display_rows"]), "kpis": s["kpis"]}
                for s in sheets_summary
            ],
        }

    # 3. PDF Generation
    try:
        import pymupdf as fitz
        from core.converter.pdf_toolkit import _find_system_font

        pdf_doc = fitz.open()
        ttf_font = _find_system_font()
        for sdata in sheets_summary:
            page = pdf_doc.new_page(width=595, height=842)
            fontname = "helv"
            if ttf_font:
                try:
                    page.insert_font(fontname="vn_font", fontfile=ttf_font)
                    fontname = "vn_font"
                except Exception:
                    pass
            y_cursor = 42.0
            page.insert_text((40, y_cursor), report_title, fontname=fontname, fontsize=13)
            y_cursor += 22.0
            page.insert_text((40, y_cursor), f"Sheet: {sdata['sheet_name']}", fontname=fontname, fontsize=11)
            y_cursor += 18.0

            for k in sdata["kpis"]:
                line = (
                    f"KPI [{k['column']}]: SUM={_format_cell_display(k['sum'])} | "
                    f"AVG={_format_cell_display(k['average'])} | COUNT={k['count']}"
                )
                page.insert_text((40, y_cursor), line, fontname=fontname, fontsize=9)
                y_cursor += 14.0

            y_cursor += 8.0
            headers_line = " | ".join(sdata["headers"])
            page.insert_text((40, y_cursor), headers_line, fontname=fontname, fontsize=9.5)
            y_cursor += 14.0
            for drow in sdata["display_rows"]:
                if y_cursor > 790:
                    page = pdf_doc.new_page(width=595, height=842)
                    if ttf_font:
                        try:
                            page.insert_font(fontname="vn_font", fontfile=ttf_font)
                        except Exception:
                            pass
                    y_cursor = 42.0
                page.insert_text((40, y_cursor), " | ".join(drow), fontname=fontname, fontsize=9)
                y_cursor += 13.0

        if len(pdf_doc) == 0:
            pdf_doc.new_page(width=595, height=842)
        pdf_doc.save(dest, garbage=3, deflate=True)
        pdf_doc.close()
    finally:
        if os.path.isfile(docx_target) and docx_target != dest:
            try:
                os.remove(docx_target)
            except OSError:
                pass

    return {
        "success": True,
        "output_path": dest,
        "format": "pdf",
        "sheets_processed": len(sheets_summary),
        "sheets": [
            {"sheet_name": s["sheet_name"], "rows": len(s["display_rows"]), "kpis": s["kpis"]}
            for s in sheets_summary
        ],
    }


def consolidate_workbooks(
    input_paths: List[str],
    output_path: str,
    mode: str = "append_rows",
    strict_schema: bool = True,
) -> Dict[str, Any]:
    """
    Consolidates multiple Excel workbooks into a single Master Workbook.
    - mode="append_rows": Validates column headers across all files and merges rows into
      a unified sheet with a 'Source_Workbook' provenance column, leading-zero '@' format
      protection (ERR_XLSX_007), and live dynamic summary formulas.
    - mode="merge_sheets": Copies worksheets from each source file into a single Master
      Workbook with an Index_Dashboard navigation sheet.
    """
    if not input_paths:
        raise ValueError("input_paths cannot be empty")

    dest = os.path.abspath(output_path)
    os.makedirs(os.path.dirname(dest), exist_ok=True)

    master_wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="0F294A", end_color="0F294A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    if mode == "merge_sheets":
        idx_ws = master_wb.active
        idx_ws.title = "Index_Dashboard"
        idx_ws["A1"] = "MASTER CONSOLIDATED WORKBOOK — INDEX DASHBOARD"
        idx_ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F294A")
        idx_ws.append(["STT", "Tên Sheet", "Tệp Nguồn", "Số Dòng", "Liên Kết Nhanh"])
        for c_idx in range(1, 6):
            cell = idx_ws.cell(row=2, column=c_idx)
            cell.fill = header_fill
            cell.font = header_font

        sheet_counter = 0
        for p in input_paths:
            abs_p = os.path.abspath(p)
            if not os.path.isfile(abs_p):
                raise FileNotFoundError(f"Input workbook not found: {abs_p}")
            src_name = os.path.splitext(os.path.basename(abs_p))[0]
            wb_v = openpyxl.load_workbook(abs_p, data_only=True)
            wb_f = openpyxl.load_workbook(abs_p, data_only=False)
            try:
                for sname in wb_f.sheetnames:
                    ws_f = wb_f[sname]
                    ws_v = wb_v[sname]
                    extracted = _extract_sheet_matrix_and_kpis(ws_v, ws_f)
                    if not extracted["headers"]:
                        continue
                    sheet_counter += 1
                    safe_title = f"{src_name[:14]}_{sname[:14]}"[:31]
                    if safe_title in master_wb.sheetnames:
                        safe_title = f"{safe_title[:27]}_{sheet_counter}"
                    new_ws = master_wb.create_sheet(title=safe_title)
                    new_ws.append(extracted["headers"])
                    for r_vals in extracted["data_rows"]:
                        new_ws.append(r_vals)
                    apply_worksheet_print_setup(new_ws)
                    idx_ws.append([
                        sheet_counter,
                        safe_title,
                        os.path.basename(abs_p),
                        len(extracted["data_rows"]),
                        f'=HYPERLINK("#\'{safe_title}\'!A1", "Mở {safe_title}")',
                    ])
            finally:
                wb_v.close()
                wb_f.close()

        apply_worksheet_print_setup(idx_ws)
        master_wb.save(dest)
        master_wb.close()
        return {
            "success": True,
            "output_path": dest,
            "mode": "merge_sheets",
            "workbooks_merged": len(input_paths),
            "sheets_created": sheet_counter,
        }

    # Default: mode="append_rows"
    ws = master_wb.active
    ws.title = "Consolidated_Data"
    canonical_headers: Optional[List[str]] = None
    all_rows_merged = 0
    schema_mismatches: List[str] = []

    for p in input_paths:
        abs_p = os.path.abspath(p)
        if not os.path.isfile(abs_p):
            raise FileNotFoundError(f"Input workbook not found: {abs_p}")
        src_basename = os.path.basename(abs_p)
        wb_v = openpyxl.load_workbook(abs_p, data_only=True)
        wb_f = openpyxl.load_workbook(abs_p, data_only=False)
        try:
            ws_f = wb_f.active
            ws_v = wb_v.active
            extracted = _extract_sheet_matrix_and_kpis(ws_v, ws_f)
            headers = extracted["headers"]
            if not headers:
                continue

            if canonical_headers is None:
                canonical_headers = headers
                full_headers = ["Source_Workbook"] + canonical_headers
                ws.append(full_headers)
                for col_i in range(1, len(full_headers) + 1):
                    c_obj = ws.cell(row=1, column=col_i)
                    c_obj.fill = header_fill
                    c_obj.font = header_font
                    c_obj.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                norm_curr = [h.strip().lower() for h in headers]
                norm_base = [h.strip().lower() for h in canonical_headers]
                if norm_curr != norm_base:
                    msg = f"Schema mismatch in '{src_basename}': expected {canonical_headers}, got {headers}"
                    schema_mismatches.append(msg)
                    if strict_schema:
                        raise ValueError(msg)

            for r_vals in extracted["data_rows"]:
                first_str = str(r_vals[0] or "").strip().lower() if r_vals else ""
                if any(kw in first_str for kw in ("tổng cộng", "total", "grand total")):
                    continue
                row_idx = ws.max_row + 1
                ws.cell(row=row_idx, column=1, value=src_basename).border = cell_border
                for col_offset, val in enumerate(r_vals[: len(canonical_headers)], start=2):
                    cell = ws.cell(row=row_idx, column=col_offset)
                    if isinstance(val, str) and val.startswith("0") and len(val) > 1 and val.isdigit():
                        cell.number_format = "@"
                        cell.value = val
                    else:
                        cell.value = val
                    cell.border = cell_border
                all_rows_merged += 1
        finally:
            wb_v.close()
            wb_f.close()

    # Add live dynamic summary row at the bottom
    if canonical_headers and all_rows_merged > 0:
        total_row_idx = ws.max_row + 1
        last_data_row = total_row_idx - 1
        lbl_cell = ws.cell(row=total_row_idx, column=1, value="TỔNG CỘNG (TOTAL)")
        lbl_cell.font = Font(name="Calibri", size=11, bold=True)
        lbl_cell.border = cell_border

        for col_i in range(2, len(canonical_headers) + 2):
            col_let = get_column_letter(col_i)
            sample_vals = [ws.cell(row=r, column=col_i).value for r in range(2, last_data_row + 1)]
            has_num = any(_to_number(v) is not None for v in sample_vals)
            t_cell = ws.cell(row=total_row_idx, column=col_i)
            if has_num:
                t_cell.value = f"=SUM({col_let}2:{col_let}{last_data_row})"
            else:
                t_cell.value = f"=COUNTA({col_let}2:{col_let}{last_data_row})"
            t_cell.font = Font(name="Calibri", size=11, bold=True)
            t_cell.border = cell_border

    ws.views.sheetView[0].showGridLines = True
    apply_worksheet_print_setup(ws)
    master_wb.save(dest)
    master_wb.close()

    return {
        "success": True,
        "output_path": dest,
        "mode": "append_rows",
        "workbooks_merged": len(input_paths),
        "rows_consolidated": all_rows_merged,
        "total_rows_consolidated": all_rows_merged,
        "headers": canonical_headers or [],
        "schema_mismatches": schema_mismatches,
    }
